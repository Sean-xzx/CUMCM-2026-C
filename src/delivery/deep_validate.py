"""对最终5个官方工作簿做结构、时间、数值和储能约束深度校验。"""
from __future__ import annotations
import json, math, sys
from datetime import date, datetime
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

BASE = Path(r"C:\Users\12055\Desktop\2026 国赛")
ANS = Path(r"C:\Users\12055\Desktop\答案")
DATES = [d.date() for d in __import__('pandas').date_range('2025-02-01','2025-12-31')]
RMAX = 5000/6
TOL = 1e-5


def dval(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    return __import__('pandas').to_datetime(v).date()


def finite_nonneg(vals):
    a=np.asarray(vals,float)
    return bool(np.isfinite(a).all() and (a >= -TOL).all())


def compare_headers(final, template, sheets):
    wf=load_workbook(final, read_only=True, data_only=False)
    wt=load_workbook(template, read_only=True, data_only=False)
    out={}
    for s in sheets:
        hf=[wf[s].cell(1,c).value for c in range(1,wf[s].max_column+1)]
        ht=[wt[s].cell(1,c).value for c in range(1,wt[s].max_column+1)]
        out[s]=(hf==ht)
    return out


def plan_check(ws):
    dates=[dval(ws.cell(r,1).value) for r in range(2,336)]
    vals=[]; maxerr=0.; formulas=0
    for r in range(2,336):
        row=[ws.cell(r,c).value for c in range(2,146)]
        vals.extend(row)
        total=float(ws.cell(r,146).value)
        maxerr=max(maxerr,abs(total-sum(float(x) for x in row)))
        for c in range(2,148):
            if isinstance(ws.cell(r,c).value,str) and ws.cell(r,c).value.startswith('='): formulas+=1
    return {'rows':334,'dates_exact':dates==DATES,'values_finite_nonnegative':finite_nonneg(vals),
            'daily_total_max_abs_error':maxerr,'formula_cells':formulas}


def storage_check(ws):
    rows=ws.max_row-1
    expected_intervals=['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']
    ds=[]; allvals=[]; throughput=[]; s0=[]; s24=[]; labels=[]
    for i in range(334):
        b=2+6*i
        ds.append(dval(ws.cell(b,1).value)); s0.append(float(ws.cell(b,6).value)); s24.append(float(ws.cell(b+1,6).value))
        for j in range(6):
            labels.append(ws.cell(b+j,2).value)
            c=float(ws.cell(b+j,3).value); u=float(ws.cell(b+j,4).value)
            allvals += [c,u]; throughput.append(c+u)
    return {'rows':rows,'rows_exact':rows==2004,'dates_exact':ds==DATES,
            'interval_labels_exact':labels==expected_intervals*334,
            'charge_discharge_finite_nonnegative':finite_nonneg(allvals),
            'max_4h_throughput_kWh':max(throughput),
            # 表内每行是24个十分钟槽的4小时汇总，因此上限为24*RMAX。
            'power_limit_ok':max(throughput)<=24*RMAX+TOL,
            'soc_endpoints_in_bounds':min(s0+s24)>=1200-TOL and max(s0+s24)<=10800+TOL,
            'soc_daily_continuity_max_abs_error':float(np.max(np.abs(np.array(s0[1:])-np.array(s24[:-1]))))}


def emergency_check(ws):
    dates=[]; bad_pairs=0; pattern=True
    for i in range(334):
        b=2+3*i
        dates.append(dval(ws.cell(b,1).value))
        if ws.cell(b+1,1).value not in (None,'') or ws.cell(b+2,1).value not in (None,''): pattern=False
        labs=(ws.cell(b,2).value or '').split(); amts=(ws.cell(b,3).value or '').split()
        if len(labs)!=len(amts): bad_pairs += 1
    return {'rows':ws.max_row-1,'rows_exact':ws.max_row-1==1002,'dates_exact':dates==DATES,
            'three_rows_per_date_pattern':pattern,'interval_amount_pair_mismatches':bad_pairs}


def main():
    report={}
    # result1: 保持模板标签，并按位置写入144个真实决策时段。
    r1=ANS/'result1.xlsx'; t1=BASE/'附件5'/'result1.xlsx'
    w=load_workbook(r1,data_only=False); wt=load_workbook(t1,data_only=False)
    ws=w['计划购电量']; vals=[ws.cell(r,2).value for r in range(2,146)]
    sys.path.insert(0,str(Path(r'C:\Users\12055\Desktop\第一问预测+优化')))
    from solve_q1 import load_attachment1, solve_dispatch
    solved=solve_dispatch(load_attachment1(Path(r'C:\Users\12055\Desktop\第一问预测+优化\附件1.xlsx')))
    solved_grid=solved.grid_to_load+solved.grid_to_storage
    report['result1']={'sheet_names_exact':w.sheetnames==wt.sheetnames,
                       'headers_and_time_labels_exact':[ws.cell(r,1).value for r in range(1,146)]==[wt['计划购电量'].cell(r,1).value for r in range(1,146)],
                       '144_values':len(vals)==144,'values_finite_nonnegative':finite_nonneg(vals),
                       'sequential_slot_mapping_max_abs_error':float(np.max(np.abs(np.asarray(vals,float)-solved_grid))),
                       'lp_success':bool(solved.success),'storage_interface_limit_ok':bool(np.max(solved.pv_to_storage+solved.grid_to_storage+solved.discharge)<=RMAX+TOL),
                       'soc_bounds_ok':bool((solved.soc>=1200-TOL).all() and (solved.soc<=10800+TOL).all())}
    specs={
      'result2':(['计划购电量','充放电量','紧急购电量'],BASE/'附件5'/'result2.xlsx'),
      'result3':(['计划购电量','调整购电量','充放电量','紧急购电量'],BASE/'附件5'/'result3.xlsx'),
      'result4-2':(['计划购电量','充放电量','紧急购电量'],BASE/'附件5'/'result4-2.xlsx'),
      'result4-3':(['计划购电量','调整购电量','充放电量','紧急购电量'],BASE/'附件5'/'result4-3.xlsx')}
    for stem,(sheets,tmpl) in specs.items():
        p=ANS/f'{stem}.xlsx'; wb=load_workbook(p,data_only=False)
        item={'sheet_names_exact':wb.sheetnames==sheets,'headers_exact':compare_headers(p,tmpl,sheets)}
        item['计划购电量']=plan_check(wb['计划购电量'])
        if '调整购电量' in sheets: item['调整购电量']=plan_check(wb['调整购电量'])
        item['充放电量']=storage_check(wb['充放电量'])
        item['紧急购电量']=emergency_check(wb['紧急购电量'])
        report[stem]=item
    ok=True
    def walk(x):
        nonlocal ok
        if isinstance(x,dict):
            for k,v in x.items():
                if isinstance(v, bool) and (k.endswith('_ok') or k.endswith('_exact') or k in {'lp_success','144_values','values_finite_nonnegative','charge_discharge_finite_nonnegative','power_limit_ok','soc_endpoints_in_bounds','dates_exact','interval_labels_exact','three_rows_per_date_pattern'}) and v is not True:
                    ok=False
                walk(v)
    walk(report)
    report['all_checks_passed']=ok
    out=ANS/'深度校验.json'; out.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'all_checks_passed':ok,'report':str(out)},ensure_ascii=False))
    if not ok: raise SystemExit(1)

if __name__=='__main__': main()
