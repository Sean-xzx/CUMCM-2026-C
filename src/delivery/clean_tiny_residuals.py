"""清理交付工作簿中的LP数值残差（|v|<1e-9 视为0），并同步重算全天购电量列。

只改 计划购电量/调整购电量 两张表的数值单元格与其合计列，不触碰工作表结构、
表头、样式、充放电表与紧急购电表，也不改变全天购电费列。
"""
from __future__ import annotations
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

ANS = Path(r"C:\Users\12055\Desktop\答案")
EPS = 1e-9
SPECS = {
    "result1": ["计划购电量"],
    "result2": ["计划购电量"],
    "result3": ["计划购电量", "调整购电量"],
    "result4-2": ["计划购电量"],
    "result4-3": ["计划购电量", "调整购电量"],
}


def main() -> dict:
    report = {}
    for stem, sheets in SPECS.items():
        path = ANS / f"{stem}.xlsx"
        wb = load_workbook(path)
        item = {}
        for name in sheets:
            ws = wb[name]
            fixed = 0
            maxdrift = 0.0
            for r in range(2, ws.max_row + 1):
                total = ws.cell(r, 146).value
                if total is None:
                    continue
                row_sum = 0.0
                for c in range(2, 146):
                    v = ws.cell(r, c).value
                    if v is None:
                        continue
                    v = float(v)
                    if 0 < abs(v) < EPS:
                        ws.cell(r, c).value = 0.0
                        fixed += 1
                        v = 0.0
                    row_sum += v
                maxdrift = max(maxdrift, abs(row_sum - float(total)))
                if not np.isclose(row_sum, float(total), rtol=0, atol=1e-9):
                    ws.cell(r, 146).value = float(row_sum)
            item[name] = {"tiny_zeroed": fixed, "max_total_drift_before_fix": maxdrift}
        wb.save(path)
        report[stem] = item
    (ANS / "残差清理报告.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == "__main__":
    main()
