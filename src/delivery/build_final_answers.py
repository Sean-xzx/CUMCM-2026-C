"""生成并核验2026国赛C题的5个官方结果文件与题目指定汇总表。"""
from __future__ import annotations

import json
import math
import shutil
import sys
from collections import defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(r"C:\Users\12055\Desktop")
OFFICIAL = ROOT / "2026 国赛"
TEMPLATES = OFFICIAL / "附件5"
ANSWER = ROOT / "答案"
Q2_DIR = ROOT / "第二问预测&优化"
Q3_SOURCE = ROOT / "第三问预测+优化" / "result3_filled.xlsx"
Q42_SOURCE = ROOT / "第四问预测+优化" / "EW方案结果" / "result4-2_filled.xlsx"
Q43_SOURCE = ROOT / "第四问预测+优化" / "EW方案结果" / "result4-3_filled.xlsx"
START = 31
T = 144
DT = 1 / 6
DATES = pd.date_range("2025-02-01", "2025-12-31", freq="D")
SPEC_DATES = pd.to_datetime(["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"])
SLOT_INDEX = [60, 72, 84, 96, 108, 120]
SLOT_LABELS = ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"]
EPS = 1e-6


def as_date(value):
    if value is None:
        return None
    try:
        return pd.Timestamp(value).normalize()
    except Exception:
        return None


def interval_label(a: int, b: int) -> str:
    def f(m: int) -> str:
        return "24:00" if m == 1440 else f"{m // 60}:{m % 60:02d}"
    return f"{f(a * 10)}-{f(b * 10)}"


def contiguous_emergency(values: np.ndarray) -> tuple[str, str, float]:
    values = np.maximum(np.asarray(values, float), 0.0)
    active = np.flatnonzero(values > EPS)
    if not len(active):
        return "", "", 0.0
    groups = []
    start = prev = int(active[0])
    for j in map(int, active[1:]):
        if j == prev + 1:
            prev = j
        else:
            groups.append((start, prev + 1))
            start = prev = j
    groups.append((start, prev + 1))
    labels = [interval_label(a, b) for a, b in groups]
    amounts = [float(values[a:b].sum()) for a, b in groups]
    return " ".join(labels), " ".join(f"{x:.6f}" for x in amounts), float(sum(amounts))


def fill_dense_plan_sheet(ws, dates, arr, price, fee_mode="plan", qplan=None):
    if arr.shape != (334, 144):
        raise ValueError(f"计划矩阵维度错误：{arr.shape}")
    for i, day in enumerate(dates, start=2):
        ws.cell(i, 1).value = day.to_pydatetime()
        ws.cell(i, 1).number_format = "yyyy/mm/dd"
        vals = np.maximum(arr[i - 2], 0.0)
        for j, value in enumerate(vals, start=2):
            ws.cell(i, j).value = float(value)
            ws.cell(i, j).number_format = "0.000000"
        ws.cell(i, 146).value = float(vals.sum())
        if fee_mode == "plan":
            fee = np.dot(price, vals)
        elif fee_mode == "adjusted":
            qp = np.asarray(qplan[i - 2], float)
            fee = np.dot(price, qp) + np.dot(1.5 * price, np.maximum(vals - qp, 0)) - np.dot(0.5 * price, np.maximum(qp - vals, 0))
        else:
            raise ValueError(fee_mode)
        ws.cell(i, 147).value = float(fee)
        ws.cell(i, 146).number_format = ws.cell(i, 147).number_format = "0.000000"


def fill_storage_sheet(ws, dates, charge, discharge, s0, soc):
    ws.delete_rows(2, max(0, ws.max_row - 1))
    row = 2
    for i, day in enumerate(dates):
        for g in range(6):
            sl = slice(g * 24, (g + 1) * 24)
            if g == 0:
                ws.cell(row, 1).value = day.to_pydatetime()
                ws.cell(row, 1).number_format = "yyyy/mm/dd"
                ws.cell(row, 5).value = "0:00"
                ws.cell(row, 6).value = float(s0[i])
            elif g == 1:
                ws.cell(row, 5).value = "24:00"
                ws.cell(row, 6).value = float(soc[i, -1])
            ws.cell(row, 2).value = f"{g * 4}:00-{(g + 1) * 4}:00"
            ws.cell(row, 3).value = float(np.maximum(charge[i, sl], 0).sum())
            ws.cell(row, 4).value = float(np.maximum(discharge[i, sl], 0).sum())
            for c in (3, 4, 6):
                ws.cell(row, c).number_format = "0.000000"
            row += 1


def fill_emergency_sheet(ws, dates, emergency):
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    ws.delete_rows(2, max(0, ws.max_row - 1))
    # 严格保留模板的“每个日期3行”结构；首行按题面表4把同日多个
    # 连续区间及其区间总量分别用空格连接，另两行保留为空。
    row = 2
    for i, day in enumerate(dates):
        labels, amounts, _ = contiguous_emergency(emergency[i])
        ws.cell(row, 1).value = day.to_pydatetime()
        ws.cell(row, 1).number_format = "yyyy/mm/dd"
        ws.cell(row, 2).value = labels or None
        ws.cell(row, 3).value = amounts or None
        for rr in range(3):
            for cc in range(1, 4):
                ws.cell(row + rr, cc)._style = copy(style_rows[rr][cc - 1])
        # 样式复制可能把日期格式覆盖为 General，必须在最后重新指定。
        ws.cell(row, 1).number_format = "yyyy/mm/dd"
        # 空字符串在Excel中显示为空，但可确保最后一个日期块的两条续行
        # 被真实保存，从而严格得到模板要求的1002条数据行。
        ws.cell(row + 1, 1).value = ""
        ws.cell(row + 2, 1).value = ""
        row += 3


def build_result2() -> dict:
    sys.path.insert(0, str(Q2_DIR))
    import q2_pipeline as q2

    dates, load, pv = q2.load_project_data(Q2_DIR / "附件2_训练集_2025年1月.csv", Q2_DIR / "附件2_预测集_2025年2月起.csv")
    price = pd.read_csv(Q2_DIR / "附件1_分时电价.csv", encoding="utf-8-sig")["电价"].to_numpy(float)
    fl1, fl2 = q2.rolling_ew_forecasts(load, dates, "load", start=14, end=len(dates), n=8, decay=0.6)
    fg1, fg2 = q2.rolling_ew_forecasts(pv, dates, "pv", start=14, end=len(dates), n=8, decay=0.6)
    net = (load - pv) * DT
    residual = net - (fl1 - fg1) * DT

    qrows, charge, discharge, emergency, soc, s0rows = [], [], [], [], [], []
    s = 6000.0
    price48 = np.r_[price, price]
    statuses = []
    for d in range(START, len(dates)):
        rho = np.quantile(residual[d - 14:d], 0.8, axis=0)
        need48 = np.r_[fl1[d] * DT + rho, fl2[d] * DT + rho]
        pv48 = np.r_[fg1[d], fg2[d]] * DT
        q, status = q2.solve_plan_lp(need48, pv48, price48, s, periods_per_day=T, eta=0.9, rmax=5000 / 6, smin=1200, smax=10800)
        if not status["success"]:
            raise RuntimeError(status["message"])
        out = q2.settle_day(q, net[d], s, eta=0.9, rmax=5000 / 6, smin=1200, smax=10800)
        s0rows.append(s)
        qrows.append(q)
        charge.append(out["charge"])
        discharge.append(out["discharge"])
        emergency.append(out["emergency"])
        soc.append(out["soc"])
        statuses.append(status["status"])
        s = float(out["soc"][-1])

    qarr = np.asarray(qrows); carr = np.asarray(charge); uarr = np.asarray(discharge)
    earr = np.asarray(emergency); sarr = np.asarray(soc); s0arr = np.asarray(s0rows)
    out_path = ANSWER / "result2.xlsx"
    wb = load_workbook(TEMPLATES / "result2.xlsx")
    fill_dense_plan_sheet(wb["计划购电量"], DATES, qarr, price)
    fill_storage_sheet(wb["充放电量"], DATES, carr, uarr, s0arr, sarr)
    fill_emergency_sheet(wb["紧急购电量"], DATES, earr)
    wb.save(out_path)
    return {"q": qarr, "charge": carr, "discharge": uarr, "emergency": earr, "soc": sarr, "s0": s0arr,
            "price": price, "lp_all_success": all(x == 0 for x in statuses), "path": str(out_path)}


def read_emergency_matrix(ws) -> np.ndarray:
    out = np.zeros((334, 144), float)
    date_to_i = {d.normalize(): i for i, d in enumerate(DATES)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        day = as_date(row[0])
        if day not in date_to_i or not row[1] or row[2] is None:
            continue
        labels = str(row[1]).split()
        amounts = str(row[2]).split()
        if len(labels) != len(amounts):
            # 旧导出一般是一行一个10分钟槽；若无法配对则直接报错。
            raise ValueError(f"紧急购电区间与数量无法配对：{row}")
        for label, amount in zip(labels, amounts):
            start_text, end_text = label.split("-")
            h, m = map(int, start_text.split(":"))
            j = (h * 60 + m) // 10
            end_h, end_m = map(int, end_text.replace("+1", "").split(":"))
            end_min = 1440 if (end_h == 0 and "+1" in end_text) or end_h == 24 else end_h * 60 + end_m
            k = end_min // 10
            total = float(amount)
            if not (0 <= j < k <= 144):
                raise ValueError(f"非法紧急购电区间：{label}")
            out[date_to_i[day], j:k] += total / (k - j)
    return out


def copy_and_normalize(source: Path, output_name: str) -> dict:
    if not source.exists():
        raise FileNotFoundError(source)
    target = ANSWER / output_name
    shutil.copy2(source, target)
    wb = load_workbook(target)
    emergency = read_emergency_matrix(wb["紧急购电量"])
    fill_emergency_sheet(wb["紧急购电量"], DATES, emergency)
    wb.save(target)
    return {"path": str(target), "emergency": emergency}


def workbook_summary(path: Path, question: str):
    wb = load_workbook(path, data_only=True)
    plans = {}
    for sn in ("计划购电量", "调整购电量"):
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        rows = {}
        for r in range(2, 336):
            day = as_date(ws.cell(r, 1).value)
            if day is None:
                continue
            vals = np.array([ws.cell(r, c).value for c in range(2, 146)], float)
            rows[day] = {"vals": vals, "total": float(ws.cell(r, 146).value), "fee": float(ws.cell(r, 147).value)}
        plans[sn] = rows

    storage = defaultdict(list)
    if "充放电量" in wb.sheetnames:
        ws = wb["充放电量"]; current = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                current = as_date(row[0])
            if current is not None and row[1] is not None:
                storage[current].append(row)

    emerg = {}
    if "紧急购电量" in wb.sheetnames:
        ws = wb["紧急购电量"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            day = as_date(row[0])
            if day is not None:
                emerg[day] = (row[1] or "", row[2] or "")
    return {"question": question, "plans": plans, "storage": storage, "emergency": emerg}


def build_required_tables(q1_summary: dict, books: list[dict], output: Path):
    wb = Workbook(); wb.remove(wb.active)
    blue = PatternFill("solid", fgColor="1F4E78"); white = Font(color="FFFFFF", bold=True)

    ws = wb.create_sheet("表1_购电量")
    headers = ["问题", "策略", "日期"] + SLOT_LABELS + ["全天购电量/kWh", "全天购电费/元"]
    ws.append(headers)
    q1_slots = {x["时间段"]: x["购电量_kWh"] for x in q1_summary["specified_intervals"]}
    ws.append(["问题1", "计划", "典型日"] + [q1_slots[x] for x in SLOT_LABELS] + [q1_summary["total_grid_purchase_kwh"], q1_summary["total_cost_yuan"]])
    for book in books:
        for plan_name, by_day in book["plans"].items():
            strategy = "0:00计划" if plan_name == "计划购电量" else "日内调整后"
            for day in SPEC_DATES:
                item = by_day[day.normalize()]
                ws.append([book["question"], strategy, day.to_pydatetime()] + [float(item["vals"][j]) for j in SLOT_INDEX] + [item["total"], item["fee"]])

    ws2 = wb.create_sheet("表2_充放电")
    headers2 = ["问题", "日期"]
    for g in range(6): headers2 += [f"{g*4}:00-{(g+1)*4}:00充电/kWh", f"{g*4}:00-{(g+1)*4}:00放电/kWh"]
    headers2 += ["0:00储电量/kWh", "24:00储电量/kWh"]
    ws2.append(headers2)
    q1_groups = q1_summary["four_hour_summary"]
    ws2.append(["问题1", "典型日"] + [x for g in q1_groups for x in (g["充电量_kWh"], g["放电量_kWh"])] + [q1_summary["soc_initial_kwh"], q1_summary["soc_terminal_kwh"]])
    for book in books:
        for day in SPEC_DATES:
            rows = book["storage"][day.normalize()]
            if len(rows) != 6: raise ValueError(f"{book['question']} {day.date()} 充放电行数不是6")
            vals = [x for row in rows for x in (float(row[2]), float(row[3]))]
            s0 = next(float(r[5]) for r in rows if str(r[4]) == "0:00")
            s24 = next(float(r[5]) for r in rows if str(r[4]) == "24:00")
            ws2.append([book["question"], day.to_pydatetime()] + vals + [s0, s24])

    ws3 = wb.create_sheet("表3_紧急购电")
    ws3.append(["问题", "日期", "紧急购电时间段", "各连续时段紧急购电量/kWh"])
    for book in books:
        for day in SPEC_DATES:
            interval, amount = book["emergency"].get(day.normalize(), ("", ""))
            ws3.append([book["question"], day.to_pydatetime(), interval, amount])

    note = wb.create_sheet("填写说明", 0)
    notes = [
        ["项目", "说明"],
        ["时间槽错位", "附件数据的00:10是00:00-00:10时段终点；官方结果模板首列表头错后一槽。所有result文件保留模板表头，数值严格按第i个决策时段写入第i个数据位置。"],
        ["日期范围", "result2、result3、result4-2、result4-3均覆盖2025/2/1—2025/12/31，共334天；中间日期没有省略。"],
        ["计划表结构", "A列日期，B:EO为144个时段，EP为全天购电量，EQ为全天购电费，共147列。"],
        ["功率口径", "问题1按论文式(10)：储能外部接口侧充电量与放电量之和不超过5000/6 kWh/10分钟；效率通过SOC状态方程计入。"],
        ["紧急购电格式", "严格保留模板每个日期3行的结构；日期块首行填写，同一天多个连续区间及对应区间总量分别以空格分隔，另两行留空，与题面表4示例一致。"],
    ]
    for row in notes: note.append(row)

    for sh in wb.worksheets:
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for cell in sh[1]: cell.fill = blue; cell.font = white; cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)): cell.number_format = "0.000000"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in sh.columns:
            letter = col[0].column_letter
            sh.column_dimensions[letter].width = min(42, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(output)


def validate_result(path: Path, stem: str) -> dict:
    wb = load_workbook(path, data_only=True, read_only=False)
    expected = {
        "result1": ["计划购电量", "充放电量"],
        "result2": ["计划购电量", "充放电量", "紧急购电量"],
        "result3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
        "result4-2": ["计划购电量", "充放电量", "紧急购电量"],
        "result4-3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
    }[stem]
    report = {"sheets_ok": wb.sheetnames == expected, "size_bytes": path.stat().st_size}
    if stem == "result1":
        ws = wb["计划购电量"]
        vals = [ws.cell(r, 2).value for r in range(2, 146)]
        report.update({"periods": len(vals), "required_nonempty": all(v is not None and math.isfinite(float(v)) for v in vals), "nonnegative": min(map(float, vals)) >= -EPS})
    else:
        for sn in [x for x in ("计划购电量", "调整购电量") if x in wb.sheetnames]:
            ws = wb[sn]
            dates = [as_date(ws.cell(r, 1).value) for r in range(2, 336)]
            blank = 0; neg = 0; max_sum_error = 0.0
            for r in range(2, 336):
                vals = [ws.cell(r, c).value for c in range(2, 146)]
                blank += sum(v is None or not math.isfinite(float(v)) for v in vals)
                neg += sum(float(v) < -EPS for v in vals if v is not None)
                max_sum_error = max(max_sum_error, abs(sum(map(float, vals)) - float(ws.cell(r, 146).value)))
            report[sn] = {"dates_ok": dates == list(DATES), "required_blank": blank, "negative": neg, "max_total_error": max_sum_error}
        sws = wb["充放电量"]
        report["storage_rows"] = sws.max_row - 1
        report["storage_expected_rows"] = 334 * 6
        ews = wb["紧急购电量"]
        report["emergency_rows"] = ews.max_row - 1
        report["emergency_expected_rows"] = 334 * 3
    return report


def main():
    ANSWER.mkdir(parents=True, exist_ok=True)
    q1_json = ROOT / "第一问预测+优化" / "result1_验证.json"
    if not (ANSWER / "result1.xlsx").exists() or not q1_json.exists():
        raise FileNotFoundError("请先运行第一问预测+优化/build_result1.py")
    q1_summary = json.loads((OFFICIAL / "第一问" / "问题1_关键指标.json").read_text(encoding="utf-8"))

    q2 = build_result2()
    q3 = copy_and_normalize(Q3_SOURCE, "result3.xlsx")
    q42 = copy_and_normalize(Q42_SOURCE, "result4-2.xlsx")
    q43 = copy_and_normalize(Q43_SOURCE, "result4-3.xlsx")

    books = [
        workbook_summary(ANSWER / "result2.xlsx", "问题2"),
        workbook_summary(ANSWER / "result3.xlsx", "问题3"),
        workbook_summary(ANSWER / "result4-2.xlsx", "问题4-2"),
        workbook_summary(ANSWER / "result4-3.xlsx", "问题4-3"),
    ]
    build_required_tables(q1_summary, books, ANSWER / "题目要求汇总表.xlsx")

    validation = {stem: validate_result(ANSWER / f"{stem}.xlsx", stem) for stem in ["result1", "result2", "result3", "result4-2", "result4-3"]}
    validation["q2_lp_all_success"] = q2["lp_all_success"]
    validation["q2_soc_bounds"] = bool(np.all((q2["soc"] >= 1200 - EPS) & (q2["soc"] <= 10800 + EPS)))
    validation["q2_soc_continuity"] = bool(np.allclose(q2["s0"][1:], q2["soc"][:-1, -1]))
    validation["required_tables_exists"] = (ANSWER / "题目要求汇总表.xlsx").exists()
    for stem, report in validation.items():
        if not isinstance(report, dict):
            continue
        if not report.get("sheets_ok", True):
            raise AssertionError(f"{stem} 工作表结构错误")
        if "emergency_rows" in report and report["emergency_rows"] != report["emergency_expected_rows"]:
            raise AssertionError(f"{stem} 紧急购电表行数错误：{report['emergency_rows']}")
        if "storage_rows" in report and report["storage_rows"] != report["storage_expected_rows"]:
            raise AssertionError(f"{stem} 充放电表行数错误：{report['storage_rows']}")
    (ANSWER / "交付验证.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
