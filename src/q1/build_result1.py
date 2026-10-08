"""按官方模板填写问题1 result1.xlsx，并核验时间槽与约束。"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

from solve_q1 import load_attachment1, solve_dispatch, validate_solution, build_summary


def fill_result1(input_xlsx: Path, template_xlsx: Path, output_xlsx: Path) -> dict:
    data = load_attachment1(input_xlsx)
    sol = solve_dispatch(data)
    checks = validate_solution(data, sol)
    summary = build_summary(data, sol, checks)

    wb = load_workbook(template_xlsx)
    if wb.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError(f"模板工作表不符：{wb.sheetnames}")

    # 附件1的时间是时段终点；模板时间段整体错后一槽。
    # 按第i个决策时段顺序写入第i个模板数据位置，绝不按错误标签匹配。
    ws = wb["计划购电量"]
    grid = sol.grid_to_load + sol.grid_to_storage
    if len(grid) != 144 or ws.max_row != 145:
        raise ValueError("问题1必须恰有144个十分钟决策时段")
    for i, value in enumerate(grid, start=2):
        ws.cell(i, 2).value = float(max(value, 0.0))
        ws.cell(i, 2).number_format = "0.000000"

    ws = wb["充放电量"]
    charge = sol.pv_to_storage + sol.grid_to_storage
    for g in range(6):
        sl = slice(g * 24, (g + 1) * 24)
        ws.cell(g + 2, 2).value = float(charge[sl].sum())
        ws.cell(g + 2, 3).value = float(sol.discharge[sl].sum())
        ws.cell(g + 2, 2).number_format = "0.000000"
        ws.cell(g + 2, 3).number_format = "0.000000"
    ws.cell(2, 5).value = 6000.0
    ws.cell(3, 5).value = float(sol.soc[-1])
    ws.cell(2, 5).number_format = ws.cell(3, 5).number_format = "0.000000"

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)

    # 回读精确目标，避免仅凭save成功即宣称完成。
    out = load_workbook(output_xlsx, data_only=True)
    plan_vals = [out["计划购电量"].cell(r, 2).value for r in range(2, 146)]
    if any(v is None for v in plan_vals):
        raise AssertionError("result1计划购电量存在空值")
    if abs(sum(plan_vals) - summary["total_grid_purchase_kwh"]) > 1e-6:
        raise AssertionError("result1购电量合计与求解结果不一致")
    if abs(out["充放电量"]["E2"].value - 6000.0) > 1e-9 or abs(out["充放电量"]["E3"].value - 6000.0) > 1e-6:
        raise AssertionError("result1日初/日末SOC不一致")

    report = {
        "output": str(output_xlsx),
        "solver_status": summary["solver_status"],
        "periods": 144,
        "total_purchase_kWh": summary["total_grid_purchase_kwh"],
        "total_cost_yuan": summary["total_cost_yuan"],
        "max_constraint_residual_kWh": max(checks.values()),
        "time_slot_rule": "第i个终点标签对应第i个实际决策时段；模板错位标签不改，数值按位置写入",
        "power_limit_assumption": "按论文式(10)，储能外部接口侧充电量与放电量之和每10分钟不超过5000/6 kWh",
    }
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=Path(r"C:\Users\12055\Desktop\2026 国赛\附件1.xlsx"))
    parser.add_argument("--template", type=Path, default=Path(r"C:\Users\12055\Desktop\2026 国赛\附件5\result1.xlsx"))
    parser.add_argument("--output", type=Path, default=Path(r"C:\Users\12055\Desktop\答案\result1.xlsx"))
    args = parser.parse_args()
    report = fill_result1(args.input, args.template, args.output)
    report_path = Path(__file__).with_name("result1_验证.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
