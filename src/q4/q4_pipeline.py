from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
import json
import shutil
import time

import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy.sparse import csr_matrix, lil_matrix
from scipy.stats import pearsonr, spearmanr
from sklearn.linear_model import Ridge
from openpyxl import load_workbook

T = 144
DT = 1 / 6
START = 31
SELECT_DAYS = range(13, 31)  # 1月14—31日
ISSUE = (0, 6, 12, 18)
STAGE_SLOT = (0, 36, 72, 108)
ETA, R_MAX, SMIN, SMAX = 0.9, 5000 / 6, 1200.0, 10800.0
PEN, UP, DN = 5.0, 1.5, 0.5
EW_N, EW_DECAY, RESID_LOOK, FUSER_LOOK = 8, 0.6, 14, 14
DAY0 = pd.Timestamp("2025-01-01")
PRICE_GAMMAS = (0.0, 0.25, 0.5, 0.75, 1.0)


@dataclass
class DataBundle:
    dates: pd.DatetimeIndex
    load: np.ndarray
    pv: np.ndarray
    price: np.ndarray
    forecast_hourly: np.ndarray


def load_official_data(base: str | Path) -> DataBundle:
    base = Path(base)
    ld = pd.read_excel(base / "附件2.xlsx", sheet_name="小区负载")
    dates = pd.DatetimeIndex(pd.to_datetime(ld.iloc[:, 0]))
    load = ld.iloc[:, 1:].to_numpy(float)
    pv = pd.read_excel(base / "附件2.xlsx", sheet_name="光伏发电实际功率").iloc[:, 1:].to_numpy(float)
    price = pd.read_excel(base / "附件4.xlsx").iloc[:, 1:].to_numpy(float)
    f = pd.read_excel(base / "附件3.xlsx")
    f["日期"] = pd.to_datetime(f["日期"].ffill())
    f["h0"] = f["预报时刻"].astype(str).str.split(":").str[0].astype(int)
    fc = np.full((len(dates), 4, 25), np.nan)
    for _, row in f.iterrows():
        d = int((row["日期"] - DAY0).days)
        if 0 <= d < len(dates):
            k = ISSUE.index(int(row["h0"]))
            for h in range(1, 25):
                fc[d, k, h] = float(row[f"预报{h}小时"])
    if load.shape != (365, T) or pv.shape != load.shape or price.shape != load.shape:
        raise ValueError(f"官方附件维度异常：L={load.shape}, G={pv.shape}, C={price.shape}")
    if not np.isfinite(load).all() or not np.isfinite(pv).all() or not np.isfinite(price).all():
        raise ValueError("附件2或附件4存在非数值/缺失值")
    return DataBundle(dates, load, pv, price, fc)


def low_load_day(date) -> bool:
    return pd.Timestamp(date).dayofweek in (4, 5)


def _weights(n: int, decay: float = EW_DECAY) -> np.ndarray:
    w = decay ** np.arange(n)
    return w / w.sum()


def ew_day(y: np.ndarray, dates, cutoff: int, target_day: int, kind: str) -> np.ndarray:
    dates = pd.DatetimeIndex(dates)
    target_date = dates[target_day] if target_day < len(dates) else dates[-1] + pd.Timedelta(days=target_day-len(dates)+1)
    if kind == "load":
        typ = low_load_day(target_date)
        idx = [j for j in range(cutoff-1, -1, -1) if low_load_day(dates[j]) == typ][:EW_N]
    elif kind == "pv":
        idx = list(range(cutoff-1, max(-1, cutoff-EW_N-1), -1))
    else:
        raise ValueError(kind)
    pred = np.sum(_weights(len(idx))[:, None] * y[idx], axis=0)
    pred = np.maximum(pred, 0.0)
    if kind == "pv":
        pred[np.max(y[:cutoff], axis=0) <= 0] = 0.0
    return pred


def ew_net48(data: DataBundle, d: int) -> np.ndarray:
    n1 = (ew_day(data.load, data.dates, d, d, "load") - ew_day(data.pv, data.dates, d, d, "pv")) * DT
    n2 = (ew_day(data.load, data.dates, d, d+1, "load") - ew_day(data.pv, data.dates, d, d+1, "pv")) * DT
    return np.r_[n1, n2]


def _default_historical_net_hat(net: np.ndarray, day: int) -> np.ndarray:
    """仅供通用接口/测试使用；正式流水线传入Q2/Q3预测器。"""
    return np.r_[net[day-1], net[day-1]]


def _fit_price_correction(net, price, cutoff_day, start_slot, forecast_provider,
                          max_history_days=42):
    """用同类历史预测误差拟合 StandardScaler + Ridge(alpha=1)。"""
    slots=net.shape[1]; cutoff_abs=cutoff_day*slots+start_slot
    nf,pf=net.ravel(),price.ravel(); xs=[]; ys=[]
    for j in range(max(7,cutoff_day-max_history_days),cutoff_day):
        pred_j=np.asarray(forecast_provider(j),float)
        if pred_j.shape!=(2*slots,): raise ValueError("历史净负荷预测必须为严格48小时")
        target=j*slots+start_slot+np.arange(2*slots)
        valid=(target<cutoff_abs)&(target<len(nf)); lag=target[valid]-7*slots
        xs.append(pred_j[valid]-nf[lag]); ys.append(pf[target[valid]]-pf[lag])
    x=np.concatenate(xs)[:,None]; y=np.concatenate(ys)
    mean=x.mean(axis=0); scale=x.std(axis=0); scale[scale<1e-12]=1.0
    model=Ridge(alpha=1.0).fit((x-mean)/scale,y)
    return model,mean,scale,len(y)


def coupled_price_forecast(net: np.ndarray, price: np.ndarray, dates,
                           cutoff_day: int, net_hat48: np.ndarray,
                           gamma: float = 1.0, max_history_days: int = 42,
                           start_slot: int = 0, forecast_provider=None) -> tuple[np.ndarray, dict]:
    """单向耦合：7日前价格基线 + gamma×Ridge(预测净负荷变化)。"""
    net=np.asarray(net,float); price=np.asarray(price,float); slots=net.shape[1]
    net_hat48=np.asarray(net_hat48,float)
    if net_hat48.shape!=(2*slots,): raise ValueError("价格预测必须接收严格48小时净负荷")
    provider=forecast_provider or (lambda j:_default_historical_net_hat(net,j))
    model,mean,scale,n_train=_fit_price_correction(net,price,cutoff_day,start_slot,provider,max_history_days)
    target=cutoff_day*slots+start_slot+np.arange(2*slots); lag=target-7*slots
    if lag.min()<0 or lag.max()>=cutoff_day*slots: raise ValueError("7日基线超出决策时刻前的合法历史")
    xcur=(net_hat48-net.ravel()[lag])[:,None]
    pred=np.maximum(price.ravel()[lag]+float(gamma)*model.predict((xcur-mean)/scale),1e-6)
    dates=pd.DatetimeIndex(dates)
    meta={"gamma":float(gamma),"ridge_alpha":1.0,"n_train":int(n_train),
          "coef_scaled":float(model.coef_[0]),"intercept":float(model.intercept_),
          "history_start":str(dates[max(7,cutoff_day-max_history_days)].date()),
          "history_end":str(dates[cutoff_day-1].date())}
    return pred,meta


def select_price_gamma(net: np.ndarray, price: np.ndarray, dates, forecast_provider,
                       candidates=PRICE_GAMMAS) -> tuple[pd.DataFrame,float]:
    rows=[]
    for gamma in candidates:
        err=[]
        for d in SELECT_DAYS:
            pred,_=coupled_price_forecast(net,price,dates,d,forecast_provider(d),gamma=gamma,
                                          forecast_provider=forecast_provider)
            err.extend(pred[:net.shape[1]]-price[d])
        rows.append({"gamma":float(gamma),"一月14—31日MAE":float(np.mean(np.abs(err)))})
    table=pd.DataFrame(rows)
    best=float(table.loc[table["一月14—31日MAE"].idxmin(),"gamma"])
    return table,best


def price_metrics(data: DataBundle, gamma: float) -> pd.DataFrame:
    net=(data.load-data.pv)*DT; provider=lambda d:ew_net48(data,d); rows=[]
    for label,days in (("一月14—31日",range(13,31)),("2—12月",range(31,365))):
        lag,cp=[],[]
        for d in days:
            pred,_=coupled_price_forecast(net,data.price,data.dates,d,provider(d),gamma=gamma,
                                          forecast_provider=provider)
            lag.extend(data.price[d-7]-data.price[d]); cp.extend(pred[:T]-data.price[d])
        lag,cp=np.asarray(lag),np.asarray(cp)
        rows.append({"评价期":label,"lag7_MAE":np.abs(lag).mean(),"耦合预测_MAE":np.abs(cp).mean(),
                     "耦合预测_RMSE":np.sqrt(np.mean(cp**2))})
    return pd.DataFrame(rows)


def _controls(dates, slots: int) -> np.ndarray:
    dates = pd.DatetimeIndex(dates)
    n = len(dates)
    slot = np.tile(np.arange(slots), n)
    slot_dm = np.eye(slots, dtype=float)[slot][:, 1:]
    dow = np.repeat(dates.dayofweek.to_numpy(), slots)
    dow_dm = np.eye(7, dtype=float)[dow][:, 1:]
    doy = np.repeat(dates.dayofyear.to_numpy(), slots)
    trend = np.repeat(np.arange(n, dtype=float), slots) / max(n-1, 1)
    return np.column_stack([np.ones(n*slots), slot_dm, dow_dm,
                            np.sin(2*np.pi*doy/365), np.cos(2*np.pi*doy/365), trend])


def _residualize(y: np.ndarray, x: np.ndarray) -> np.ndarray:
    coef, *_ = np.linalg.lstsq(x, y.ravel(), rcond=None)
    return (y.ravel() - x @ coef).reshape(y.shape)


def partial_price_netload_dependence(net: np.ndarray, price: np.ndarray, dates,
                                     bootstrap: int = 400, seed: int = 42) -> dict:
    net, price = np.asarray(net, float), np.asarray(price, float)
    x = _controls(dates, net.shape[1])
    rn, rp = _residualize(net, x), _residualize(price, x)
    raw_p = pearsonr(net.ravel(), price.ravel()).statistic
    raw_s = spearmanr(net.ravel(), price.ravel()).statistic
    pp = pearsonr(rn.ravel(), rp.ravel()).statistic
    ps = spearmanr(rn.ravel(), rp.ravel()).statistic
    rng = np.random.default_rng(seed)
    vals = np.empty(bootstrap)
    for b in range(bootstrap):
        idx = rng.integers(0, len(net), len(net))
        vals[b] = pearsonr(rn[idx].ravel(), rp[idx].ravel()).statistic
    lo, hi = np.quantile(vals, [0.025, 0.975])
    return {"raw_pearson": float(raw_p), "raw_spearman": float(raw_s),
            "partial_pearson": float(pp), "partial_spearman": float(ps),
            "ci_low": float(lo), "ci_high": float(hi), "bootstrap_days": int(bootstrap)}


def solve_q2_plan48(load_energy, pv_energy, price_hat48, s0,
                    periods_per_day=T, eta=ETA, rmax=R_MAX, smin=SMIN, smax=SMAX):
    """问题二原LP口径：显式区分小区负荷与光伏，只替换为预测电价。"""
    load_energy=np.asarray(load_energy,float); pv_energy=np.asarray(pv_energy,float); price_hat48=np.asarray(price_hat48,float)
    h=2*periods_per_day
    if load_energy.shape!=(h,) or pv_energy.shape!=(h,) or price_hat48.shape!=(h,):
        raise ValueError("问题4-2输入必须恰为48小时")
    nv=6*h; eu,es,qu,qs,v,soc=[np.arange(h)+k*h for k in range(6)]
    c=np.zeros(nv); c[qu]=price_hat48; c[qs]=price_hat48
    bounds=[(0,None)]*(5*h)+[(smin,smax)]*h
    aub=lil_matrix((2*h,nv)); bub=np.r_[pv_energy,np.full(h,rmax)]
    for t in range(h):
        aub[t,eu[t]]=1; aub[t,es[t]]=1
        aub[h+t,es[t]]=1; aub[h+t,qs[t]]=1; aub[h+t,v[t]]=1
    aeq=lil_matrix((2*h,nv)); beq=np.r_[load_energy,s0,np.zeros(h-1)]
    for t in range(h):
        aeq[t,eu[t]]=1; aeq[t,qu[t]]=1; aeq[t,v[t]]=1
        r=h+t; aeq[r,soc[t]]=1; aeq[r,es[t]]=-eta; aeq[r,qs[t]]=-eta; aeq[r,v[t]]=1/eta
        if t>0: aeq[r,soc[t-1]]=-1
    res=linprog(c,A_ub=aub.tocsr(),b_ub=bub,A_eq=aeq.tocsr(),b_eq=beq,bounds=bounds,method="highs")
    if not res.success:
        return {"success":False,"status":int(res.status),"message":res.message,
                "q_exec":np.full(periods_per_day,np.nan),"soc48":np.full(h,np.nan)}
    return {"success":True,"status":int(res.status),"message":res.message,
            "q_exec":res.x[qu][:periods_per_day]+res.x[qs][:periods_per_day],
            "soc48":res.x[soc],"objective_predicted":float(res.fun)}


def solve_stage48(target48: np.ndarray, price_hat48: np.ndarray, s0: float,
                  periods_per_day: int = T, eta: float = ETA, rmax: float = R_MAX,
                  smin: float = SMIN, smax: float = SMAX,
                  qplan: np.ndarray | None = None, exec_slots: int | None = None) -> dict:
    """预测电价目标下的严格48小时LP；返回当前可执行部分。"""
    target48 = np.asarray(target48, float); price_hat48 = np.asarray(price_hat48, float)
    h = 2 * periods_per_day
    if target48.shape != (h,) or price_hat48.shape != (h,):
        raise ValueError("优化输入必须恰为48小时且目标与预测电价等长")
    n_exec = periods_per_day if exec_slots is None else int(exec_slots)
    nv = 4*h + (n_exec if qplan is not None else 0)
    iq, ic, iu, iS = (np.arange(h)+j*h for j in range(4))
    c = np.zeros(nv)
    if qplan is None:
        c[iq] = price_hat48
    else:
        qplan = np.asarray(qplan, float)
        if qplan.shape != (n_exec,):
            raise ValueError("qplan长度必须等于可执行时段数")
        c[iq[n_exec:]] = price_hat48[n_exec:]
        iy = 4*h + np.arange(n_exec); c[iy] = 1.0
    lb = np.zeros(nv); ub = np.full(nv, np.inf)
    lb[iS] = smin; ub[iS] = smax
    if qplan is not None: lb[iy] = -np.inf
    rows=[]; cols=[]; vals=[]; rhs=[]; rr=0; tt=np.arange(h)
    rows += [rr+tt, rr+tt, rr+tt]; cols += [iq, iu, ic]
    vals += [-np.ones(h), -np.ones(h), np.ones(h)]; rhs.append(-target48); rr += h
    rows += [rr+tt, rr+tt]; cols += [ic, iu]
    vals += [np.ones(h), np.ones(h)]; rhs.append(np.full(h, rmax)); rr += h
    if qplan is not None:
        p0 = price_hat48[:n_exec]
        for coef in (DN, UP):
            r0 = rr + np.arange(n_exec)
            rows += [r0, r0]; cols += [iq[:n_exec], iy]
            vals += [coef*p0, -np.ones(n_exec)]
            rhs.append(coef*p0*qplan); rr += n_exec
    aub = csr_matrix((np.concatenate(vals), (np.concatenate(rows), np.concatenate(cols))), shape=(rr,nv))
    er=[tt,tt[1:],tt,tt]; ec=[iS,iS[:-1],ic,iu]
    ev=[np.ones(h),-np.ones(h-1),-eta*np.ones(h),np.ones(h)/eta]
    aeq=csr_matrix((np.concatenate(ev),(np.concatenate(er),np.concatenate(ec))),shape=(h,nv))
    beq=np.zeros(h); beq[0]=s0
    res=linprog(c,A_ub=aub,b_ub=np.concatenate(rhs),A_eq=aeq,b_eq=beq,bounds=np.c_[lb,ub],method="highs")
    if not res.success:
        return {"success":False,"status":int(res.status),"message":res.message,
                "q_exec":np.full(n_exec,np.nan),"soc48":np.full(h,np.nan)}
    return {"success":True,"status":int(res.status),"message":res.message,
            "q_exec":res.x[iq[:n_exec]],"soc48":res.x[iS],"objective_predicted":float(res.fun)}


def settle_day(q: np.ndarray, net_actual: np.ndarray, s0: float,
               eta: float = ETA, rmax: float = R_MAX,
               smin: float = SMIN, smax: float = SMAX) -> dict:
    q=np.asarray(q,float); net=np.asarray(net_actual,float); n=len(q)
    charge=np.zeros(n); discharge=np.zeros(n); emergency=np.zeros(n); curtail=np.zeros(n); soc=np.zeros(n)
    s=float(s0)
    for t in range(n):
        gap=net[t]-q[t]
        if gap>0:
            discharge[t]=min(gap,rmax,(s-smin)*eta); emergency[t]=gap-discharge[t]; s-=discharge[t]/eta
        else:
            charge[t]=min(-gap,rmax,(smax-s)/eta); curtail[t]=-gap-charge[t]; s+=eta*charge[t]
        soc[t]=s
    return {"s0":float(s0),"charge":charge,"discharge":discharge,"emergency":emergency,"curtail":curtail,"soc":soc}


def state_at(q, net, s0, t0):
    return float(s0) if t0 == 0 else float(settle_day(q[:t0], net[:t0], s0)["soc"][-1])


def cost_q2(q, emergency, true_price) -> dict:
    plan=float(np.sum(np.asarray(q)*true_price)); emg=float(np.sum(PEN*np.asarray(emergency)*true_price))
    return {"计划购电费_元":plan,"紧急购电费_元":emg,"总费用_元":plan+emg}


def cost_q3(qplan, qadj, emergency, true_price) -> dict:
    qp,qa,e,p=map(np.asarray,(qplan,qadj,emergency,true_price))
    plan=float(np.sum(qp*p)); up=float(np.sum(UP*p*np.maximum(qa-qp,0)))
    dn=float(np.sum(DN*p*np.maximum(qp-qa,0))); emg=float(np.sum(PEN*p*e))
    return {"计划购电费_元":plan,"调增费_元":up,"调减退费_元":dn,
            "紧急购电费_元":emg,"总费用_元":plan+up-dn+emg}


class Q4Engine:
    """忠实承接问题二 EW/EWMA 与问题三卡尔曼四阶段口径，仅替换价格层。"""
    def __init__(self, data: DataBundle, price_gamma: float):
        self.data=data; self.gamma=float(price_gamma); self.net=(data.load-data.pv)*DT
        self.lflat=data.load.ravel(); self.gflat=data.pv.ravel(); self.nflat=self.net.ravel()
        self.fc10=np.full((len(data.dates),4,T),np.nan)
        for d in range(len(data.dates)):
            for k in range(4): self.fc10[d,k]=self.downscale(d,k)

    @lru_cache(None)
    def ew_load(self, cutoff, target): return ew_day(self.data.load,self.data.dates,cutoff,target,"load")
    @lru_cache(None)
    def ew_pv(self, cutoff, target): return ew_day(self.data.pv,self.data.dates,cutoff,target,"pv")

    def ew_abs(self, kind, cutoff, abs_slots):
        out=np.empty(len(abs_slots)); arr=self.data.load if kind=="load" else self.data.pv
        for td in np.unique(abs_slots//T):
            m=abs_slots//T==td; out[m]=ew_day(arr,self.data.dates,cutoff,int(td),kind)[abs_slots[m]%T]
        return out

    def downscale(self,d,k):
        release=d*T+STAGE_SLOT[k]; x=np.arange(release,release+T)
        anchors=release+6*np.arange(1,25)-1; vals=self.data.forecast_hourly[d,k,1:25]
        prev=self.gflat[release-1] if release>0 else 0.0
        return np.maximum(np.interp(x,np.r_[release-1,anchors],np.r_[prev,vals]),0)

    @lru_cache(None)
    def point48(self,d,k,use_attachment=True):
        release=d*T+STAGE_SLOT[k]; abs_slots=np.arange(release,release+2*T)
        lh=self.ew_abs("load",d,abs_slots); gew=self.ew_abs("pv",d,abs_slots); gh=gew.copy()
        if use_attachment:
            for r in range(T):
                efc=[]; eew=[]
                for j in range(d-1,max(7,d-FUSER_LOOK-3),-1):
                    ht=j*T+STAGE_SLOT[k]+r
                    if ht>=release or ht>=len(self.gflat): continue
                    f=self.fc10[j,k,r]; actual=self.gflat[ht]; ew0=self.ew_abs("pv",j,np.array([ht]))[0]
                    if actual>50 or f>50: efc.append(f-actual); eew.append(ew0-actual)
                    if len(efc)>=FUSER_LOOK: break
                fcur=self.fc10[d,k,r]
                if len(efc)<5: gh[r]=.5*fcur+.5*gew[r]
                else:
                    efc=np.asarray(efc); eew=np.asarray(eew); bias=efc.mean()
                    vfc=max(efc.var(),1.0); vew=max(eew.var(),1.0)
                    gh[r]=((fcur-bias)/vfc+gew[r]/vew)/(1/vfc+1/vew)
        return (lh-np.maximum(gh,0))*DT

    @lru_cache(None)
    def residual_samples48(self,d,k,look=RESID_LOOK):
        release=d*T+STAGE_SLOT[k]; samples=np.full((look,2*T),np.nan)
        for r in range(2*T):
            vals=[]
            for j in range(d-1,7,-1):
                target=j*T+STAGE_SLOT[k]+r
                if target>=release or target>=len(self.nflat): continue
                vals.append(self.nflat[target]-self.point48(j,k)[r])
                if len(vals)==look: break
            if vals: samples[:len(vals),r]=vals
        return samples

    def rho48(self,d,k,alpha):
        return np.nan_to_num(np.nanquantile(self.residual_samples48(d,k),alpha,axis=0),nan=0.0)

    def price48(self,d,k,net_hat,use_attachment=True):
        provider=lambda j:self.point48(j,k,use_attachment)
        return coupled_price_forecast(self.net,self.data.price,self.data.dates,d,net_hat,
                                      gamma=self.gamma,start_slot=STAGE_SLOT[k],
                                      forecast_provider=provider)[0]

    def simulate_q2(self, days=None, alpha=.8, s_init=6000.0):
        days=list(range(START,len(self.data.dates)) if days is None else days)
        hist=[]
        for d in range(10,days[0]): hist.append(self.net[d]-self.point48(d,0,False)[:T])
        rec={k:[] for k in ("qplan","qadj","charge","discharge","emergency","curtail","soc","s0","terminal_soc48")}; s=float(s_init); statuses=[]
        for d in days:
            load48=np.r_[self.ew_load(d,d),self.ew_load(d,d+1)]*DT
            pv48=np.r_[self.ew_pv(d,d),self.ew_pv(d,d+1)]*DT
            pred=load48-pv48
            rho=np.quantile(np.asarray(hist[-RESID_LOOK:]),alpha,axis=0)
            ph=self.price48(d,0,pred,False)
            sol=solve_q2_plan48(load48+np.r_[rho,rho],pv48,ph,s)
            if not sol["success"]: raise RuntimeError(f"Q4-2 {self.data.dates[d].date()} LP失败：{sol['message']}")
            q=sol["q_exec"]; st=settle_day(q,self.net[d],s)
            for key,val in (("qplan",q),("qadj",q),("charge",st["charge"]),("discharge",st["discharge"]),
                            ("emergency",st["emergency"]),("curtail",st["curtail"]),("soc",st["soc"])): rec[key].append(val)
            rec["s0"].append(s); rec["terminal_soc48"].append(float(sol["soc48"][-1])); s=float(st["soc"][-1]); statuses.append(sol["status"])
            hist.append(self.net[d]-pred[:T]); hist=hist[-60:]
        out={k:np.asarray(v) for k,v in rec.items()}; out["lp_status"]=np.asarray(statuses); out["days"]=np.asarray(days)
        return out

    def simulate_q3(self, days=None, a0=.6, alo=.7, ahi=.9, s_init=6000.0):
        days=list(range(START,len(self.data.dates)) if days is None else days)
        rec={k:[] for k in ("qplan","qadj","charge","discharge","emergency","curtail","soc","s0","terminal_soc48")}; statuses=[]; s=float(s_init)
        for d in days:
            pred0=self.point48(d,0); target0=pred0+self.rho48(d,0,a0); ph0=self.price48(d,0,pred0)
            sol=solve_stage48(target0,ph0,s)
            if not sol["success"]: raise RuntimeError(f"Q4-3 {self.data.dates[d].date()} 0:00 LP失败")
            q=sol["q_exec"]; qplan=q.copy(); base=target0[:T].copy(); terminal0=float(sol["soc48"][-1]); statuses.append(sol["status"])
            for k in (1,2,3):
                t0=STAGE_SLOT[k]; n=T-t0; pred=self.point48(d,k); lo=pred+self.rho48(d,k,alo); hi=pred+self.rho48(d,k,ahi)
                if np.all((base[t0:]>=lo[:n])&(base[t0:]<=hi[:n])): continue
                target=pred+self.rho48(d,k,a0); target[:n]=np.clip(base[t0:],lo[:n],hi[:n])
                ph=self.price48(d,k,pred); sn=state_at(q,self.net[d],s,t0)
                sol=solve_stage48(target,ph,sn,qplan=qplan[t0:],exec_slots=n)
                if not sol["success"]: raise RuntimeError(f"Q4-3 {self.data.dates[d].date()} {ISSUE[k]}:00 LP失败")
                q[t0:]=sol["q_exec"]; statuses.append(sol["status"])
            st=settle_day(q,self.net[d],s)
            for key,val in (("qplan",qplan),("qadj",q),("charge",st["charge"]),("discharge",st["discharge"]),
                            ("emergency",st["emergency"]),("curtail",st["curtail"]),("soc",st["soc"])): rec[key].append(val)
            rec["s0"].append(s); rec["terminal_soc48"].append(terminal0); s=float(st["soc"][-1])
        out={k:np.asarray(v) for k,v in rec.items()}; out["lp_status"]=np.asarray(statuses); out["days"]=np.asarray(days)
        return out


def ablation_table(total_costs: dict, selected_gamma: float) -> pd.DataFrame:
    """把同口径 gamma=0 与一月冻结 gamma 的平行结果整理为可审计表。"""
    rows=[]
    for scheme in ("问题4-2","问题4-3"):
        base=float(total_costs[(scheme,0.0)])
        for gamma,label in ((0.0,"lag7基线"),(float(selected_gamma),"净负荷耦合")):
            total=float(total_costs[(scheme,gamma)])
            rows.append({"方案":scheme,"gamma":gamma,"价格方案":label,"总费用_元":total,
                         "相对lag7节省_元":base-total})
    return pd.DataFrame(rows)


def _fill_common(wb, rec, dates, true_price, q3=False):
    qp,qa=rec["qplan"],rec["qadj"]
    sheets=[("计划购电量",qp)] + ([ ("调整购电量",qa) ] if q3 else [])
    for name,arr in sheets:
        ws=wb[name]
        for i,d in enumerate(dates):
            ws.cell(i+2,1).value=d.to_pydatetime()
            for j in range(T): ws.cell(i+2,j+2).value=float(max(arr[i,j],0))
            ws.cell(i+2,T+2).value=float(arr[i].sum())
            if name=="计划购电量": fee=float(np.dot(true_price[i],qp[i]))
            else:
                fee=float(np.dot(true_price[i],qp[i])+np.dot(UP*true_price[i],np.maximum(qa[i]-qp[i],0))
                          -np.dot(DN*true_price[i],np.maximum(qp[i]-qa[i],0)))
            ws.cell(i+2,T+3).value=fee
    ws=wb["充放电量"]; ws.delete_rows(2,ws.max_row); row=2
    for i,d in enumerate(dates):
        for g in range(6):
            sl=slice(g*24,(g+1)*24)
            if g==0: ws.cell(row,1).value=d.to_pydatetime(); ws.cell(row,5).value="0:00"; ws.cell(row,6).value=float(rec["s0"][i])
            elif g==1: ws.cell(row,5).value="24:00"; ws.cell(row,6).value=float(rec["soc"][i,-1])
            ws.cell(row,2).value=f"{g*4}:00-{(g+1)*4}:00"; ws.cell(row,3).value=float(rec["charge"][i,sl].sum()); ws.cell(row,4).value=float(rec["discharge"][i,sl].sum()); row+=1
    ws=wb["紧急购电量"]; ws.delete_rows(2,ws.max_row); row=2
    for i,d in enumerate(dates):
        for j in np.where(rec["emergency"][i]>1e-6)[0]:
            ws.cell(row,1).value=d.to_pydatetime(); ws.cell(row,2).value=f"{j*10//60}:{j*10%60:02d}-{(j+1)*10//60}:{(j+1)*10%60:02d}"; ws.cell(row,3).value=float(rec["emergency"][i,j]); row+=1


def export_results(rec2, rec3, data: DataBundle, template_dir: str | Path, out_dir: str | Path):
    template_dir=Path(template_dir); out_dir=Path(out_dir); dates=data.dates[START:]
    paths=[]
    for stem,rec,isq3 in (("result4-2",rec2,False),("result4-3",rec3,True)):
        wb=load_workbook(template_dir/f"{stem}.xlsx"); _fill_common(wb,rec,dates,data.price[START:],isq3)
        path=out_dir/f"{stem}_filled.xlsx"; wb.save(path)
        shutil.copy2(path,out_dir/f"{stem}.xlsx")
        paths.append(path)
    return paths


def validate_excel(template_path, filled_path, stem="") -> dict:
    wt=load_workbook(template_path,read_only=True); wf=load_workbook(filled_path,read_only=False,data_only=True)
    names=wt.sheetnames==wf.sheetnames; required=["计划购电量"]+(["调整购电量"] if "4-3" in stem else [])
    dates_ok=True; nonempty=True
    expected=pd.date_range("2025-02-01","2025-12-31")
    for name in required:
        ws=wf[name]; got=pd.DatetimeIndex(pd.to_datetime([ws.cell(i,1).value for i in range(2,336)]))
        dates_ok &= bool(np.array_equal(got.values,expected.values))
        for i in range(2,336):
            nonempty &= all(ws.cell(i,j).value is not None for j in range(2,148))
    return {"sheet_names_match":bool(names),"dates_ok":bool(dates_ok),"required_nonempty":bool(nonempty)}


def generate_figures(data, engine, rec2, rec3, pmetrics, out_dir):
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    out=Path(out_dir)/"图表"; out.mkdir(exist_ok=True)
    fonts={f.name for f in font_manager.fontManager.ttflist}; cn=next((x for x in ("Microsoft YaHei","SimHei","Noto Sans CJK SC") if x in fonts),"DejaVu Sans")
    plt.rcParams.update({"font.sans-serif":[cn,"DejaVu Sans"],"axes.unicode_minus":False,"savefig.dpi":300,"axes.spines.top":False,"axes.spines.right":False,"axes.grid":True,"grid.alpha":.25})
    # 价格预测
    d=79; prednet=engine.point48(d,0,False); ph=engine.price48(d,0,prednet)
    fig,ax=plt.subplots(figsize=(8,4),constrained_layout=True); x=np.arange(T)/6
    ax.plot(x,data.price[d],label="真实电价（仅事后）",color="#555555",ls="--"); ax.plot(x,data.price[d-7],label="7日前基线",color="#E69F00"); ax.plot(x,ph[:T],label="同类型EW耦合预测",color="#0072B2")
    ax.set(title="2025年3月21日电价预测：决策只使用同类型历史EW基线与净负荷预测",xlabel="时刻（小时）",ylabel="电价（元/kWh）",xlim=(0,24)); ax.legend(frameon=False,ncol=3)
    for ext in ("png","svg"): fig.savefig(out/f"问题四_典型日电价预测.{ext}",bbox_inches="tight",facecolor="white")
    plt.close(fig)
    # 费用分解
    c2=cost_q2(rec2["qplan"],rec2["emergency"],data.price[START:]); c3=cost_q3(rec3["qplan"],rec3["qadj"],rec3["emergency"],data.price[START:])
    labels=["问题4-2","问题4-3"]; plan=[c2["计划购电费_元"]/1e4,c3["计划购电费_元"]/1e4]; adj=[0,(c3["调增费_元"]-c3["调减退费_元"])/1e4]; emg=[c2["紧急购电费_元"]/1e4,c3["紧急购电费_元"]/1e4]
    fig,ax=plt.subplots(figsize=(6.5,4.2),constrained_layout=True); ax.bar(labels,plan,label="计划购电费",color="#0072B2"); ax.bar(labels,adj,bottom=plan,label="调整净费用",color="#E69F00"); ax.bar(labels,emg,bottom=np.array(plan)+np.array(adj),label="紧急购电费",color="#D55E00")
    ax.set(title="波动电价下问题4-2与问题4-3费用构成",ylabel="费用（万元）"); ax.legend(frameon=False)
    for ext in ("png","svg"): fig.savefig(out/f"问题四_费用构成.{ext}",bbox_inches="tight",facecolor="white")
    plt.close(fig)


def run_pipeline(base=r"C:\Users\12055\Desktop\2026 国赛", out_dir=None, bootstrap=400):
    out=Path(out_dir or Path.cwd()); out.mkdir(parents=True,exist_ok=True); t0=time.time()
    data=load_official_data(base); provider=lambda d:ew_net48(data,d)
    selection,gamma=select_price_gamma((data.load-data.pv)*DT,data.price,data.dates,provider)
    pm=price_metrics(data,gamma)
    dep_all=partial_price_netload_dependence((data.load-data.pv)*DT,data.price,data.dates,bootstrap=bootstrap)
    dep_jan=partial_price_netload_dependence((data.load[:31]-data.pv[:31])*DT,data.price[:31],data.dates[:31],bootstrap=bootstrap)
    pd.DataFrame([{"范围":"全年",**dep_all},{"范围":"一月",**dep_jan}]).to_csv(out/"偏相关指标.csv",index=False,encoding="utf-8-sig")
    selection.to_csv(out/"一月价格模型选参.csv",index=False,encoding="utf-8-sig"); pm.to_csv(out/"价格预测指标.csv",index=False,encoding="utf-8-sig")
    engine=Q4Engine(data,0.0)
    rec2_lag=engine.simulate_q2(); rec3_lag=engine.simulate_q3()
    c2_lag=cost_q2(rec2_lag["qplan"],rec2_lag["emergency"],data.price[START:])
    c3_lag=cost_q3(rec3_lag["qplan"],rec3_lag["qadj"],rec3_lag["emergency"],data.price[START:])
    engine.gamma=gamma
    rec2=engine.simulate_q2(); rec3=engine.simulate_q3()
    c2=cost_q2(rec2["qplan"],rec2["emergency"],data.price[START:]); c3=cost_q3(rec3["qplan"],rec3["qadj"],rec3["emergency"],data.price[START:])
    costs=pd.DataFrame([{"方案":"问题4-2","价格方案":"净负荷耦合",**c2,"紧急购电量_kWh":rec2["emergency"].sum(),"LP成功数":len(rec2["lp_status"])},
                        {"方案":"问题4-3","价格方案":"净负荷耦合",**c3,"紧急购电量_kWh":rec3["emergency"].sum(),"LP成功数":len(rec3["lp_status"])}])
    abl=ablation_table({("问题4-2",0.0):c2_lag["总费用_元"],("问题4-2",gamma):c2["总费用_元"],
                        ("问题4-3",0.0):c3_lag["总费用_元"],("问题4-3",gamma):c3["总费用_元"]},gamma)
    costs.to_csv(out/"问题四_费用与求解指标.csv",index=False,encoding="utf-8-sig")
    abl.to_csv(out/"问题四_gamma平行消融.csv",index=False,encoding="utf-8-sig")
    export_results(rec2,rec3,data,Path(base)/"附件5",out); generate_figures(data,engine,rec2,rec3,pm,out)
    checks={
        "price_gamma":gamma,"ridge_alpha":1.0,"q2_lp_all_success":bool(np.all(rec2["lp_status"]==0)),"q3_lp_all_success":bool(np.all(rec3["lp_status"]==0)),
        "q2_soc_bounds":bool(np.all((rec2["soc"]>=SMIN-1e-6)&(rec2["soc"]<=SMAX+1e-6))),"q3_soc_bounds":bool(np.all((rec3["soc"]>=SMIN-1e-6)&(rec3["soc"]<=SMAX+1e-6))),
        "q2_soc_continuity":bool(np.allclose(rec2["s0"][1:],rec2["soc"][:-1,-1])),"q3_soc_continuity":bool(np.allclose(rec3["s0"][1:],rec3["soc"][:-1,-1])),
        "q2_cost_identity":bool(np.isclose(c2["总费用_元"],c2["计划购电费_元"]+c2["紧急购电费_元"])),
        "q2_terminal_soc_at_lower_share":float(np.mean(np.isclose(rec2["terminal_soc48"],SMIN,atol=1e-5))),
        "q3_terminal_soc_at_lower_share":float(np.mean(np.isclose(rec3["terminal_soc48"],SMIN,atol=1e-5))),
        "q3_cost_identity":bool(np.isclose(c3["总费用_元"],c3["计划购电费_元"]+c3["调增费_元"]-c3["调减退费_元"]+c3["紧急购电费_元"])),
        "runtime_seconds":time.time()-t0,
    }
    for stem in ("result4-2","result4-3"): checks[stem]=validate_excel(Path(base)/"附件5"/f"{stem}.xlsx",out/f"{stem}_filled.xlsx",stem)
    (out/"运行验证.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding="utf-8")
    return {"data":data,"engine":engine,"price_selection":selection,"price_metrics":pm,"dependence_all":dep_all,"dependence_jan":dep_jan,"rec2":rec2,"rec3":rec3,"rec2_lag":rec2_lag,"rec3_lag":rec3_lag,"costs":costs,"ablation":abl,"checks":checks}


if __name__ == "__main__":
    result=run_pipeline(out_dir=Path(__file__).resolve().parent)
    print(result["price_metrics"].to_string(index=False))
    print(result["costs"].to_string(index=False))
    print(json.dumps(result["checks"],ensure_ascii=False,indent=2))
