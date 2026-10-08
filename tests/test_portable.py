import json
from pathlib import Path
import subprocess
import sys
import numpy as np
import pytest
import run

ROOT=Path(__file__).resolve().parents[1]

def test_preserved_files_match_original_sha256():
    assert run.check_sources()["sha256_match"]

def test_demo_runs_real_solvers_and_exports(tmp_path):
    summary=run.demo(tmp_path)
    assert summary["solver_status"]=="Optimal"
    assert summary["periods"]==144
    assert summary["total_cost_yuan"]==pytest.approx(3206.6449614242956,abs=1e-5)
    assert summary["soc_terminal_kwh"]==pytest.approx(6000.,abs=1e-6)
    assert summary["synthetic_48h_lp_success"]
    assert (tmp_path/"demo_dispatch.xlsx").is_file()
    assert (tmp_path/"demo_dispatch.csv").is_file()
    assert json.loads((tmp_path/"demo_summary.json").read_text(encoding="utf-8"))["data_kind"].startswith("synthetic")

def test_missing_resources_fail_clearly_without_writing(tmp_path):
    with pytest.raises(ValueError,match="missing="):
        run.check_resources(tmp_path)
    assert not list(tmp_path.iterdir())

def test_cli_missing_resources_returns_nonzero(tmp_path):
    result=subprocess.run([sys.executable,str(ROOT/"scripts/run.py"),"check-resources","--data-dir",str(tmp_path)],capture_output=True,text=True)
    assert result.returncode==2
    assert "Official resources missing=" in result.stderr

def test_committed_reference_workbooks_have_complete_nonnegative_purchases():
    from openpyxl import load_workbook
    for stem in ["result1","result2","result3","result4-2","result4-3"]:
        wb=load_workbook(ROOT/"results/reference/final"/f"{stem}.xlsx",data_only=True,read_only=True)
        rows=list(wb["计划购电量"].iter_rows(min_row=2,max_row=145 if stem=="result1" else 335,max_col=2 if stem=="result1" else 145,values_only=True))
        matrix=np.asarray([row[1:] for row in rows],float)
        assert matrix.shape==((144,1) if stem=="result1" else (334,144))
        assert np.isfinite(matrix).all() and (matrix>=0).all()
        wb.close()


def test_verify_rejects_out_of_bounds_exported_soc(tmp_path):
    import shutil
    from openpyxl import load_workbook
    for stem in ('result1','result2'):
        shutil.copyfile(ROOT/'results/reference/final'/f'{stem}.xlsx',tmp_path/f'{stem}.xlsx')
    wb=load_workbook(tmp_path/'result2.xlsx')
    wb['充放电量']['F2']=0.0
    wb.save(tmp_path/'result2.xlsx');wb.close()
    with pytest.raises(ValueError,match='Invalid storage: result2/soc_endpoints_in_bounds'):
        run.verify(tmp_path)
