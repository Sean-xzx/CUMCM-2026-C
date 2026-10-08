"""Portable entry points around byte-for-byte preserved competition sources."""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import sys
import time

for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
for directory in ("q1", "q2", "q4", "delivery"):
    sys.path.insert(0, str(ROOT / "src" / directory))


def save_json(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def load_source(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def check_sources():
    entries = json.loads((ROOT / "source-manifest.json").read_text(encoding="utf-8"))
    for entry in entries:
        path = ROOT / entry["path"]
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != entry["sha256"]:
            raise ValueError(f"Preserved file changed or missing: {entry['path']}")
    return {"preserved_files": len(entries), "sha256_match": True}


def check_resources(data_dir):
    base = Path(data_dir).resolve()
    resources = json.loads((ROOT / "data/resource-manifest.json").read_text(encoding="utf-8"))["files"]
    missing, changed = [], []
    for entry in resources:
        path = base / entry["path"]
        if not path.is_file():
            missing.append(entry["path"])
        elif hashlib.sha256(path.read_bytes()).hexdigest() != entry["sha256"]:
            changed.append(entry["path"])
    if missing or changed:
        raise ValueError(f"Official resources missing={missing}, SHA-256 mismatch={changed}. See docs/RESOURCES.md.")
    return base


@contextmanager
def working_directory(path):
    old = Path.cwd()
    Path(path).mkdir(parents=True, exist_ok=True)
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(old)


def demo(out):
    import numpy as np
    import solve_q1 as q1
    import q4_pipeline as q4
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    slots = np.arange(144)
    hours = (slots + 1) / 6
    load = np.full(144, 600.0)
    pv = np.maximum(900 * np.sin(np.pi * (hours - 6) / 12), 0)
    price = np.where((hours > 8) & (hours <= 20), 1.2, .4)
    labels = [f"{int(h):02d}:{int(round((h % 1) * 60)):02d}" for h in hours[:-1]] + ["0:00+1"]
    data = q1.DispatchData(labels, price, load, pv, load / 6, pv / 6)
    solution = q1.solve_dispatch(data)
    if not solution.success:
        raise RuntimeError(solution.message)
    checks = q1.validate_solution(data, solution)
    summary = q1.build_summary(data, solution, checks)
    q1.export_workbook(out / "demo_dispatch.xlsx", data, solution, summary)
    import pandas as pd
    pd.DataFrame({"slot":slots,"purchase_kwh":solution.grid_to_load+solution.grid_to_storage,"soc_kwh":solution.soc}).to_csv(out / "demo_dispatch.csv",index=False)
    # A real shared 48-hour optimization + realized settlement, without official data.
    target = np.full(288, 100.)
    plan = q4.solve_stage48(target, np.tile(price,2), 6000.)
    if not plan["success"]:
        raise RuntimeError("Synthetic 48-hour optimization failed")
    actual = target[:144] + np.sin(slots / 10) * 5
    settled = q4.settle_day(plan["q_exec"], actual, 6000.)
    assert np.all((settled["soc"] >= 1200-1e-6) & (settled["soc"] <= 10800+1e-6))
    summary["synthetic_48h_lp_success"] = True
    summary["data_kind"] = "synthetic; not the competition result"
    save_json(out / "demo_summary.json", summary)
    print(json.dumps({k:summary[k] for k in ["solver_status","periods","total_cost_yuan","total_grid_purchase_kwh","soc_terminal_kwh","synthetic_48h_lp_success"]},indent=2))
    return summary


def prepare_q2_inputs(base, destination):
    """Derive original CSV schema from hash-checked Excel; never write original inputs."""
    import numpy as np
    import pandas as pd
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    load = pd.read_excel(base / "附件2.xlsx", sheet_name="小区负载").iloc[:,1:].to_numpy(float)
    pv = pd.read_excel(base / "附件2.xlsx", sheet_name="光伏发电实际功率").iloc[:,1:].to_numpy(float)
    timestamps = pd.date_range("2025-01-01 00:10", periods=365*144, freq="10min")
    frame = pd.DataFrame({"时间":timestamps,"小区负载":load.ravel(),"光伏发电实际功率":pv.ravel()})
    split = timestamps < pd.Timestamp("2025-02-01")
    for name,mask in [("附件2_训练集_2025年1月.csv",split),("附件2_预测集_2025年2月起.csv",~split)]:
        frame.loc[mask].to_csv(destination/name,index=False,encoding="utf-8-sig")
    price = pd.read_excel(base / "附件1.xlsx")["电价"].to_numpy(float)
    pd.DataFrame({"时间":np.arange(144),"电价":price}).to_csv(destination/"附件1_分时电价.csv",index=False,encoding="utf-8-sig")


def reproduce_q1(base, out):
    from build_result1 import fill_result1
    report = fill_result1(base/"附件1.xlsx",base/"附件5/result1.xlsx",out/"result1.xlsx")
    report.pop("output",None)
    save_json(out/"q1_summary.json",report)
    return report


def reproduce_q2(base, out):
    delivery = load_source("portable_delivery",ROOT/"src/delivery/build_final_answers.py")
    prepare_q2_inputs(base,out/"_q2_inputs")
    delivery.ANSWER = out
    delivery.TEMPLATES = base/"附件5"
    delivery.Q2_DIR = out/"_q2_inputs"
    record = delivery.build_result2()
    import numpy as np
    price = record["price"][None,:]
    total = float((record["q"]*price).sum()+5*(record["emergency"]*price).sum())
    report = {"total_cost_yuan":total,"days":334,"lp_all_success":record["lp_all_success"]}
    save_json(out/"q2_summary.json",report)
    return report


def reproduce_q3(base, out, select=False):
    # Execute the unchanged source cells. Relative resources are staged in an ignored
    # work directory, so no rewriting or monkey-patching of the model is necessary.
    nb = json.loads((ROOT/"src/q3/问题三_预测与优化.ipynb").read_text(encoding="utf-8"))
    work = out/"_q3_work"
    work.mkdir(parents=True,exist_ok=True)
    for name in ["附件1.xlsx","附件2.xlsx","附件3.xlsx"]:
        shutil.copyfile(base/name,work/name)
    shutil.copyfile(base/"附件5/result3.xlsx",work/"result3.xlsx")
    namespace = {"__name__":"portable_q3","display":lambda *a,**k:None}
    with working_directory(work):
        exec(compile("".join(nb["cells"][2]["source"]),"original_q3_cell2","exec"),namespace)
        if select:
            _,params = namespace["select_parameters_january"]()
        else:
            params = {"A0":.6,"ALO":.7,"AHI":.9}
        record = namespace["simulate"](stages=(1,2,3),**params)
        cost = namespace["cost"](record)
        namespace.update(FIN=record,PARAMS=params)
        for cell in (16,18):
            exec(compile("".join(nb["cells"][cell]["source"]),f"original_q3_cell{cell}","exec"),namespace)
    delivery = load_source("portable_q3_delivery",ROOT/"src/delivery/build_final_answers.py")
    delivery.ANSWER=out
    delivery.copy_and_normalize(work/"result3_filled.xlsx","result3.xlsx")
    report={"cost":cost,"total_cost_yuan":float(cost["总费用/万元"])*10000,"parameters":params,"january_parameter_selection_rerun":select}
    save_json(out/"q3_summary.json",report)
    return report


def reproduce_q4(base,out):
    import q4_ew_compare as ew
    # This unchanged entry point uses the archived old-scheme cost CSV only as its
    # comparison baseline. It independently recomputes both final EW simulations.
    shutil.copyfile(ROOT/"results/reference/q4/问题四_费用与求解指标.csv",out/"问题四_费用与求解指标.csv")
    _,_,checks=ew.run_compare(base=base,out_dir=out)
    delivery=load_source("portable_q4_delivery",ROOT/"src/delivery/build_final_answers.py")
    delivery.ANSWER=out
    for stem in ["result4-2","result4-3"]:
        delivery.copy_and_normalize(out/"EW方案结果"/f"{stem}_filled.xlsx",f"{stem}.xlsx")
    save_json(out/"q4_summary.json",checks)
    return checks


def verify(results):
    import numpy as np
    from openpyxl import load_workbook
    results=Path(results).resolve()
    reference=ROOT/"results/reference/final"
    report={}
    deep=load_source("portable_result_checks",ROOT/"src/delivery/deep_validate.py")
    for stem in ["result1","result2","result3","result4-2","result4-3"]:
        path=results/f"{stem}.xlsx"
        if not path.exists():
            raise FileNotFoundError(f"Missing generated result: {stem}.xlsx")
        actual=load_workbook(path,data_only=True,read_only=True)
        expected=load_workbook(reference/f"{stem}.xlsx",data_only=True,read_only=True)
        if actual.sheetnames!=expected.sheetnames:
            raise ValueError(f"Sheet mismatch: {stem}")
        diffs={}
        for name in ["计划购电量"]+(["调整购电量"] if "调整购电量" in actual.sheetnames else []):
            stop=2 if stem=="result1" else 145
            def matrix(ws):
                return np.asarray([[v for v in row[1:stop]] for row in ws.iter_rows(min_row=2,max_row=145 if stem=="result1" else 335,max_col=stop,values_only=True)],float)
            aa,bb=matrix(actual[name]),matrix(expected[name])
            if not np.isfinite(aa).all() or (aa < -1e-8).any():
                raise ValueError(f"Invalid purchase values: {stem}/{name}")
            delta=float(np.max(np.abs(aa-bb)))
            diffs[name]=delta
            if delta>1e-4:
                raise ValueError(f"Reference mismatch {stem}/{name}: {delta:.6g} kWh")
        # Original workbook validators inspect cell ranges repeatedly: use normal
        # mode here, rather than repeatedly streaming a read-only worksheet.
        actual.close();expected.close()
        actual=load_workbook(path,data_only=True)
        expected=load_workbook(reference/f"{stem}.xlsx",data_only=True)
        for name in actual.sheetnames:
            if [c.value for c in actual[name][1]] != [c.value for c in expected[name][1]]:
                raise ValueError(f"Header mismatch: {stem}/{name}")
        structure={"headers_match_reference":True}
        if stem=="result1":
            labels=[actual["计划购电量"].cell(r,1).value for r in range(1,146)]
            if labels != [expected["计划购电量"].cell(r,1).value for r in range(1,146)]:
                raise ValueError("Q1 template time labels changed")
            structure["time_labels_match_reference"]=True
        else:
            for name in ["计划购电量"]+(["调整购电量"] if "调整购电量" in actual.sheetnames else []):
                item=deep.plan_check(actual[name])
                if not item["dates_exact"] or not item["values_finite_nonnegative"] or item["daily_total_max_abs_error"]>1e-4 or item["formula_cells"]:
                    raise ValueError(f"Invalid plan structure/totals: {stem}/{name}")
                structure[name]=item
            storage=deep.storage_check(actual["充放电量"])
            for key in ("rows_exact","dates_exact","interval_labels_exact","charge_discharge_finite_nonnegative","power_limit_ok","soc_endpoints_in_bounds"):
                if not storage[key]:raise ValueError(f"Invalid storage: {stem}/{key}")
            if storage["soc_daily_continuity_max_abs_error"]>1e-4:
                raise ValueError(f"Broken daily SOC continuity: {stem}")
            emergency=deep.emergency_check(actual["紧急购电量"])
            if not all(emergency[k] for k in ("rows_exact","dates_exact","three_rows_per_date_pattern")) or emergency["interval_amount_pair_mismatches"]:
                raise ValueError(f"Invalid emergency export: {stem}")
            structure.update(storage=storage,emergency=emergency)
        report[stem]={"purchase_max_abs_error_kwh":diffs,"within_1e-4_kwh":True,"structure":structure}
        actual.close();expected.close()
    save_json(results/"reference-comparison.json",report)
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


def main():
    p=argparse.ArgumentParser(description=__doc__)
    sub=p.add_subparsers(dest="command",required=True)
    sub.add_parser("check-sources")
    demo_p=sub.add_parser("demo");demo_p.add_argument("--out",type=Path,default=Path("runs/demo"))
    resources=sub.add_parser("check-resources");resources.add_argument("--data-dir",type=Path,required=True)
    full=sub.add_parser("reproduce");full.add_argument("--data-dir",type=Path,required=True)
    full.add_argument("--out",type=Path,default=Path("runs/full"));full.add_argument("--questions",nargs="+",choices=["q1","q2","q3","q4"],default=["q1","q2","q3","q4"])
    full.add_argument("--select-q3",action="store_true",help="Rerun 16 January Q3 parameter combinations instead of using recorded parameters")
    v=sub.add_parser("verify");v.add_argument("--results",type=Path,default=Path("runs/full"))
    args=p.parse_args()
    check_sources()
    if args.command=="check-sources":print(json.dumps(check_sources(),indent=2))
    elif args.command=="check-resources":check_resources(args.data_dir);print("All 9 resource SHA-256 hashes match.")
    elif args.command=="demo":demo(args.out)
    elif args.command=="verify":verify(args.results)
    else:
        base=check_resources(args.data_dir);out=args.out.resolve();out.mkdir(parents=True,exist_ok=True)
        reports={};start=time.perf_counter()
        for question in args.questions:
            print(f"Running {question} ...",flush=True)
            reports[question]=reproduce_q3(base,out,args.select_q3) if question=="q3" else globals()["reproduce_"+question](base,out)
            print(f"Completed {question}",flush=True)
        save_json(out/"run-summary.json",{"questions":reports,"elapsed_seconds":time.perf_counter()-start})


if __name__=="__main__":
    try:
        main()
    except (FileNotFoundError,ValueError) as exc:
        print(f"ERROR: {exc}",file=sys.stderr)
        raise SystemExit(2)
