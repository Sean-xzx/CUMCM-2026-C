"""2026 年国赛 C 题问题 1：微网计划购电优化。"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from typing import Sequence

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import linprog
from scipy.sparse import lil_matrix


DELTA_T_HOURS = 1.0 / 6.0


@dataclass
class DispatchData:
    """问题 1 的逐时段输入数据。"""

    end_times: list[str]
    price: np.ndarray
    load_power: np.ndarray
    pv_power: np.ndarray
    load_energy: np.ndarray
    pv_energy: np.ndarray


@dataclass
class DispatchSolution:
    """线性规划返回的逐时段决策结果。"""

    success: bool
    status: int
    message: str
    total_cost: float
    pv_to_load: np.ndarray
    pv_to_storage: np.ndarray
    q_to_load: np.ndarray
    q_to_storage: np.ndarray
    discharge: np.ndarray
    soc: np.ndarray


def _normalise_end_time(value: object) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value).strip()
    if text in {"0:00+1", "00:00+1"}:
        return "0:00+1"
    if ":" in text:
        parts = text.split(":")
        try:
            return f"{int(parts[0]):02d}:{int(parts[1]):02d}"
        except (ValueError, IndexError):
            pass
    return text


def interval_labels_from_end_times(end_times: Sequence[str]) -> list[str]:
    """把表中“区间结束时刻”转换为十分钟区间标签。"""

    labels: list[str] = []
    for raw in end_times:
        text = _normalise_end_time(raw)
        if text == "0:00+1":
            end_minutes = 24 * 60
        else:
            try:
                hour_text, minute_text = text.split(":")[:2]
                end_minutes = int(hour_text) * 60 + int(minute_text)
            except (ValueError, IndexError) as exc:
                raise ValueError(f"无法识别时间：{raw}") from exc
        start_minutes = end_minutes - 10
        if start_minutes < 0:
            raise ValueError(f"结束时刻早于第一个十分钟区间：{raw}")

        def format_minutes(value: int) -> str:
            if value == 24 * 60:
                return "24:00"
            return f"{value // 60:02d}:{value % 60:02d}"

        labels.append(f"{format_minutes(start_minutes)}-{format_minutes(end_minutes)}")
    return labels


def interval_plot_coordinates(periods: int) -> tuple[np.ndarray, np.ndarray]:
    """返回十分钟区间的边界和中心坐标（小时）。"""

    if periods <= 0:
        raise ValueError("绘图区间数必须为正整数")
    edges = np.arange(periods + 1, dtype=float) * DELTA_T_HOURS
    centres = edges[:-1] + DELTA_T_HOURS / 2.0
    return edges, centres


def load_attachment1(path: str | Path) -> DispatchData:
    """读取附件 1，并把 kW 按 10 分钟换算为 kWh。"""

    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"找不到附件 1：{path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if not rows:
        raise ValueError("附件 1 没有数据行")
    if any(len(row) < 4 or any(value is None for value in row[:4]) for row in rows):
        raise ValueError("附件 1 的时间、电价、负载或光伏列存在空值")

    end_times = [_normalise_end_time(row[0]) for row in rows]
    price = np.asarray([float(row[1]) for row in rows], dtype=float)
    load_power = np.asarray([float(row[2]) for row in rows], dtype=float)
    pv_power = np.asarray([float(row[3]) for row in rows], dtype=float)

    if np.any(price < 0) or np.any(load_power < 0) or np.any(pv_power < 0):
        raise ValueError("附件 1 出现负电价、负负载或负光伏功率")

    return DispatchData(
        end_times=end_times,
        price=price,
        load_power=load_power,
        pv_power=pv_power,
        load_energy=load_power * DELTA_T_HOURS,
        pv_energy=pv_power * DELTA_T_HOURS,
    )


def solve_dispatch(
    data: DispatchData,
    *,
    eta: float = 0.9,
    rate_limit: float = 5000.0 * DELTA_T_HOURS,
    soc_min: float = 1200.0,
    soc_max: float = 10800.0,
    soc_initial: float = 6000.0,
    soc_terminal: float = 6000.0,
) -> DispatchSolution:
    """按题设线性规划求解问题 1。

    所有流量变量均为每个时段累计电量（kWh）。充电量与放电量
    之和不超过单时段吞吐上限，允许十分钟内分时切换工作模式。
    """

    n = len(data.price)
    arrays = (
        data.load_power,
        data.pv_power,
        data.load_energy,
        data.pv_energy,
    )
    if n == 0 or any(len(values) != n for values in arrays):
        raise ValueError("输入数组长度不一致或为空")
    if not 0.0 < eta <= 1.0:
        raise ValueError("eta 必须位于 (0, 1] 内")
    if rate_limit <= 0.0 or soc_min > soc_max:
        raise ValueError("充放电上限或储能容量边界无效")
    if not soc_min <= soc_initial <= soc_max:
        raise ValueError("初始储电量不在容量范围内")
    if not soc_min <= soc_terminal <= soc_max:
        raise ValueError("终止储电量不在容量范围内")

    # 变量分块：e^u, e^s, q^u, q^s, d, S。
    def variable(block: int, period: int) -> int:
        return block * n + period

    objective = np.zeros(6 * n, dtype=float)
    objective[2 * n : 3 * n] = data.price
    objective[3 * n : 4 * n] = data.price

    # 不等式统一写成 A_ub @ x <= b_ub。
    # 1) e^u + q^u + d >= l；2) e^u + e^s <= G；
    # 3) e^s + q^s + d <= R_max（区间内分时充放电）。
    a_ub = lil_matrix((3 * n, 6 * n), dtype=float)
    b_ub = np.zeros(3 * n, dtype=float)
    for t in range(n):
        a_ub[t, variable(0, t)] = -1.0
        a_ub[t, variable(2, t)] = -1.0
        a_ub[t, variable(4, t)] = -1.0
        b_ub[t] = -data.load_energy[t]

        a_ub[n + t, variable(0, t)] = 1.0
        a_ub[n + t, variable(1, t)] = 1.0
        b_ub[n + t] = data.pv_energy[t]

        a_ub[2 * n + t, variable(1, t)] = 1.0
        a_ub[2 * n + t, variable(3, t)] = 1.0
        a_ub[2 * n + t, variable(4, t)] = 1.0
        b_ub[2 * n + t] = rate_limit

    # n 条状态转移方程，加 1 条终止储电量方程。
    a_eq = lil_matrix((n + 1, 6 * n), dtype=float)
    b_eq = np.zeros(n + 1, dtype=float)
    for t in range(n):
        a_eq[t, variable(5, t)] = 1.0
        if t == 0:
            b_eq[t] = soc_initial
        else:
            a_eq[t, variable(5, t - 1)] = -1.0
        a_eq[t, variable(1, t)] = -eta
        a_eq[t, variable(3, t)] = -eta
        a_eq[t, variable(4, t)] = 1.0 / eta

    a_eq[n, variable(5, n - 1)] = 1.0
    b_eq[n] = soc_terminal

    bounds = [(0.0, None)] * (5 * n) + [(soc_min, soc_max)] * n
    result = linprog(
        objective,
        A_ub=a_ub.tocsr(),
        b_ub=b_ub,
        A_eq=a_eq.tocsr(),
        b_eq=b_eq,
        bounds=bounds,
        method="highs",
    )
    if not result.success:
        empty = np.full(n, np.nan)
        return DispatchSolution(
            success=False,
            status=int(result.status),
            message=str(result.message),
            total_cost=float("nan"),
            pv_to_load=empty.copy(),
            pv_to_storage=empty.copy(),
            q_to_load=empty.copy(),
            q_to_storage=empty.copy(),
            discharge=empty.copy(),
            soc=empty.copy(),
        )

    x = np.asarray(result.x, dtype=float)
    return DispatchSolution(
        success=True,
        status=int(result.status),
        message=str(result.message),
        total_cost=float(result.fun),
        pv_to_load=x[0:n].copy(),
        pv_to_storage=x[n : 2 * n].copy(),
        q_to_load=x[2 * n : 3 * n].copy(),
        q_to_storage=x[3 * n : 4 * n].copy(),
        discharge=x[4 * n : 5 * n].copy(),
        soc=x[5 * n : 6 * n].copy(),
    )


def validate_solution(
    data: DispatchData,
    solution: DispatchSolution,
    *,
    eta: float = 0.9,
    rate_limit: float = 5000.0 * DELTA_T_HOURS,
    soc_min: float = 1200.0,
    soc_max: float = 10800.0,
    soc_initial: float = 6000.0,
    soc_terminal: float = 6000.0,
    tolerance: float = 1e-6,
) -> dict[str, float]:
    """独立重算所有核心约束；发现违反时立即报错。"""

    if not solution.success:
        raise ValueError(f"求解器未返回可行最优解：{solution.message}")

    flows = (
        solution.pv_to_load,
        solution.pv_to_storage,
        solution.q_to_load,
        solution.q_to_storage,
        solution.discharge,
    )
    n = len(data.price)
    if any(len(values) != n for values in (*flows, solution.soc)):
        raise ValueError("解向量长度与输入时段数不一致")
    if not np.isfinite(solution.total_cost) or any(
        not np.all(np.isfinite(values)) for values in (*flows, solution.soc)
    ):
        raise ValueError("目标值和全部解向量必须为有限数")
    if min(float(np.min(values)) for values in flows) < -tolerance:
        raise ValueError("解中出现负的电量决策")

    charge = solution.pv_to_storage + solution.q_to_storage
    throughput = charge + solution.discharge
    if float(np.max(throughput - rate_limit)) > tolerance:
        raise ValueError("充放电总量超过单时段上限")

    pv_used = solution.pv_to_load + solution.pv_to_storage
    if float(np.max(pv_used - data.pv_energy)) > tolerance:
        raise ValueError("使用的光伏电量超过可用光伏电量")

    supply = solution.pv_to_load + solution.q_to_load + solution.discharge
    shortage = data.load_energy - supply
    if float(np.max(shortage)) > tolerance:
        raise ValueError("微网供电量低于小区负载")

    if float(np.min(solution.soc)) < soc_min - tolerance:
        raise ValueError("储电量低于下限")
    if float(np.max(solution.soc)) > soc_max + tolerance:
        raise ValueError("储电量高于上限")

    previous_soc = np.r_[soc_initial, solution.soc[:-1]]
    expected_soc = previous_soc + eta * charge - solution.discharge / eta
    state_residual = solution.soc - expected_soc
    if float(np.max(np.abs(state_residual))) > tolerance:
        raise ValueError("储能状态转移方程不成立")
    if abs(float(solution.soc[-1]) - soc_terminal) > tolerance:
        raise ValueError("终止储电量不符合要求")

    recalculated_cost = float(
        np.dot(data.price, solution.q_to_load + solution.q_to_storage)
    )
    if abs(recalculated_cost - solution.total_cost) > max(tolerance, tolerance * abs(recalculated_cost)):
        raise ValueError("目标函数值与逐时段费用不一致")

    return {
        "max_shortage_kwh": max(0.0, float(np.max(shortage))),
        "max_supply_surplus_kwh": max(0.0, float(np.max(-shortage))),
        "max_pv_excess_kwh": max(0.0, float(np.max(pv_used - data.pv_energy))),
        "max_throughput_excess_kwh": max(0.0, float(np.max(throughput - rate_limit))),
        "max_state_residual_kwh": float(np.max(np.abs(state_residual))),
        "terminal_soc_residual_kwh": abs(float(solution.soc[-1]) - soc_terminal),
    }


def build_summary(
    data: DispatchData,
    solution: DispatchSolution,
    validation: dict[str, float],
    *,
    soc_initial: float = 6000.0,
) -> dict[str, object]:
    """汇总论文中最常用的购电、储能与校验指标。"""

    labels = interval_labels_from_end_times(data.end_times)
    q = solution.q_to_load + solution.q_to_storage
    charge = solution.pv_to_storage + solution.q_to_storage
    pv_used = solution.pv_to_load + solution.pv_to_storage
    curtailed = np.maximum(data.pv_energy - pv_used, 0.0)
    interval_cost = data.price * q
    baseline_q = np.maximum(data.load_energy - data.pv_energy, 0.0)
    baseline_cost = float(np.dot(data.price, baseline_q))
    savings = baseline_cost - solution.total_cost
    savings_rate = savings / baseline_cost if baseline_cost > 0 else 0.0

    specified_labels = [
        "10:00-10:10",
        "12:00-12:10",
        "14:00-14:10",
        "16:00-16:10",
        "18:00-18:10",
        "20:00-20:10",
    ]
    label_to_index = {label: index for index, label in enumerate(labels)}
    specified_intervals = []
    for label in specified_labels:
        index = label_to_index[label]
        specified_intervals.append(
            {
                "时间段": label,
                "购电量_kWh": float(q[index]),
                "时段购电费_元": float(interval_cost[index]),
            }
        )

    four_hour_summary = []
    periods_per_group = 4 * 6
    for group in range(6):
        start = group * periods_per_group
        end = (group + 1) * periods_per_group
        four_hour_summary.append(
            {
                "时间段": f"{group * 4:02d}:00-{(group + 1) * 4:02d}:00",
                "充电量_kWh": float(np.sum(charge[start:end])),
                "放电量_kWh": float(np.sum(solution.discharge[start:end])),
            }
        )

    return {
        "solver_status": "Optimal" if solution.success else "Failed",
        "solver_message": solution.message,
        "periods": len(data.price),
        "total_cost_yuan": float(solution.total_cost),
        "total_purchase_q_kwh": float(np.sum(q)),
        "q_to_load_kwh": float(np.sum(solution.q_to_load)),
        "q_to_storage_kwh": float(np.sum(solution.q_to_storage)),
        "total_load_kwh": float(np.sum(data.load_energy)),
        "total_pv_available_kwh": float(np.sum(data.pv_energy)),
        "total_pv_used_kwh": float(np.sum(pv_used)),
        "pv_curtailed_kwh": float(np.sum(curtailed)),
        "total_charge_kwh": float(np.sum(charge)),
        "total_discharge_kwh": float(np.sum(solution.discharge)),
        "soc_initial_kwh": float(soc_initial),
        "soc_terminal_kwh": float(solution.soc[-1]),
        "soc_min_observed_kwh": float(np.min(solution.soc)),
        "soc_max_observed_kwh": float(np.max(solution.soc)),
        "no_storage_baseline_cost_yuan": baseline_cost,
        "cost_saving_vs_no_storage_yuan": float(savings),
        "cost_saving_rate": float(savings_rate),
        "intervals_with_both_charge_and_discharge": int(
            np.sum((charge > 1e-7) & (solution.discharge > 1e-7))
        ),
        "specified_intervals": specified_intervals,
        "four_hour_summary": four_hour_summary,
        "validation": validation,
    }


def _set_chinese_plot_style() -> str:
    available = {font.name for font in font_manager.fontManager.ttflist}
    candidates = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
    ]
    selected = next((font for font in candidates if font in available), "DejaVu Sans")
    plt.rcParams.update(
        {
            "font.sans-serif": [selected, "DejaVu Sans"],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.edgecolor": "#333333",
            "axes.grid": True,
            "grid.alpha": 0.28,
            "grid.linewidth": 0.6,
            "savefig.dpi": 300,
        }
    )
    return selected


def export_workbook(
    path: Path,
    data: DispatchData,
    solution: DispatchSolution,
    summary: dict[str, object],
) -> None:
    """导出逐时段结果、关键指标和四小时汇总。"""

    labels = interval_labels_from_end_times(data.end_times)
    charge = solution.pv_to_storage + solution.q_to_storage
    q = solution.q_to_load + solution.q_to_storage
    pv_used = solution.pv_to_load + solution.pv_to_storage
    interval_cost = data.price * q
    curtailed = np.maximum(data.pv_energy - pv_used, 0.0)
    supply_margin = (
        solution.pv_to_load
        + solution.q_to_load
        + solution.discharge
        - data.load_energy
    )

    workbook = Workbook()
    detail = workbook.active
    detail.title = "逐时段结果"
    headers = [
        "序号",
        "时间段",
        "结束时刻",
        "电价(元/kWh)",
        "负载功率(kW)",
        "光伏功率(kW)",
        "负载电量(kWh)",
        "可用光伏电量(kWh)",
        "光伏直供(kWh)",
        "光伏充电(kWh)",
        "外购直供(kWh)",
        "外购充电(kWh)",
        "储能放电(kWh)",
        "总充电量(kWh)",
        "总购电量(kWh)",
        "购电费(元)",
        "时段末储电量(kWh)",
        "弃光量(kWh)",
        "供电裕量(kWh)",
    ]
    detail.append(headers)
    for t in range(len(data.price)):
        row = [
            t + 1,
            labels[t],
            data.end_times[t],
            data.price[t],
            data.load_power[t],
            data.pv_power[t],
            data.load_energy[t],
            data.pv_energy[t],
            solution.pv_to_load[t],
            solution.pv_to_storage[t],
            solution.q_to_load[t],
            solution.q_to_storage[t],
            solution.discharge[t],
            charge[t],
            q[t],
            interval_cost[t],
            solution.soc[t],
            curtailed[t],
            supply_margin[t],
        ]
        detail.append([float(value) if isinstance(value, np.floating) else value for value in row])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    detail.column_dimensions["A"].width = 8
    detail.column_dimensions["B"].width = 19
    detail.column_dimensions["C"].width = 12
    for column in "DEFGHIJKLMNOPQRS":
        detail.column_dimensions[column].width = 18
    for row in detail.iter_rows(min_row=2, min_col=4):
        for cell in row:
            cell.number_format = "0.0000"

    metrics = workbook.create_sheet("关键指标")
    metrics.append(["指标", "数值", "单位"])
    metric_rows = [
        ("全天购电量", summary["total_purchase_q_kwh"], "kWh"),
        ("全天购电费", summary["total_cost_yuan"], "元"),
        ("全天负载电量", summary["total_load_kwh"], "kWh"),
        ("光伏可用电量", summary["total_pv_available_kwh"], "kWh"),
        ("弃光量", summary["pv_curtailed_kwh"], "kWh"),
        ("总充电量", summary["total_charge_kwh"], "kWh"),
        ("总放电量", summary["total_discharge_kwh"], "kWh"),
        ("无储能基线购电费", summary["no_storage_baseline_cost_yuan"], "元"),
        ("相对基线节省", summary["cost_saving_vs_no_storage_yuan"], "元"),
        ("相对基线节省率", summary["cost_saving_rate"], "比例"),
        ("最低储电量", summary["soc_min_observed_kwh"], "kWh"),
        ("最高储电量", summary["soc_max_observed_kwh"], "kWh"),
    ]
    for row in metric_rows:
        metrics.append(row)

    selected = workbook.create_sheet("指定时段")
    selected.append(["时间段", "购电量(kWh)", "时段购电费(元)"])
    for item in summary["specified_intervals"]:
        selected.append([item["时间段"], item["购电量_kWh"], item["时段购电费_元"]])

    grouped = workbook.create_sheet("四小时汇总")
    grouped.append(["时间段", "充电量(kWh)", "放电量(kWh)"])
    for item in summary["four_hour_summary"]:
        grouped.append([item["时间段"], item["充电量_kWh"], item["放电量_kWh"]])

    for sheet in (metrics, selected, grouped):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 20
        for row in sheet.iter_rows(min_row=2, min_col=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0000"

    workbook.save(path)


def plot_overview(
    path: Path,
    data: DispatchData,
    solution: DispatchSolution,
    summary: dict[str, object],
) -> None:
    """生成包含输入、购电、储能和校验结果的总览图。"""

    _set_chinese_plot_style()
    n = len(data.price)
    interval_edges, interval_centres = interval_plot_coordinates(n)
    q = solution.q_to_load + solution.q_to_storage
    charge = solution.pv_to_storage + solution.q_to_storage

    figure, axes = plt.subplot_mosaic(
        [["power", "price"], ["q", "battery"], ["soc", "summary"]],
        figsize=(14, 10),
        layout="constrained",
    )
    figure.get_layout_engine().set(rect=(0.0, 0.0, 1.0, 0.93))
    figure.suptitle(
        "问题1：微网计划购电与储能调度结果",
        y=0.995,
        fontsize=17,
        fontweight="bold",
    )
    figure.text(
        0.5,
        0.965,
        "数据来源：2026年国赛C题附件1；功率按每10分钟区间平均值换算为电量。",
        ha="center",
        va="top",
        fontsize=8.5,
        color="#555555",
    )

    axes["power"].plot(
        interval_centres,
        data.load_power,
        label="小区负载",
        color="#D55E00",
    )
    axes["power"].plot(
        interval_centres,
        data.pv_power,
        label="光伏预测功率",
        color="#E69F00",
    )
    axes["power"].set_title("(a) 负载与光伏预测功率")
    axes["power"].set_ylabel("功率（kW）")
    axes["power"].legend(frameon=False)

    axes["price"].stairs(data.price, interval_edges, color="#7A3E9D", baseline=None)
    axes["price"].set_title("(b) 计划电价")
    axes["price"].set_ylabel("电价（元/kWh）")

    axes["q"].stairs(
        q,
        interval_edges,
        color="#0072B2",
        fill=True,
        alpha=0.28,
    )
    axes["q"].stairs(
        q,
        interval_edges,
        color="#0072B2",
        label="总购电量",
    )
    axes["q"].stairs(
        solution.q_to_storage,
        interval_edges,
        color="#56B4E9",
        linestyle="--",
        label="其中：购电充能",
    )
    axes["q"].set_title("(c) 十分钟计划购电量")
    axes["q"].set_ylabel("电量（kWh/10 min）")
    axes["q"].legend(frameon=False)

    width = DELTA_T_HOURS * 0.86
    axes["battery"].bar(
        interval_centres,
        charge,
        width=width,
        color="#009E73",
        label="充电",
    )
    axes["battery"].bar(
        interval_centres,
        -solution.discharge,
        width=width,
        color="#CC79A7",
        label="放电（负向显示）",
    )
    axes["battery"].axhline(0.0, color="#333333", linewidth=0.8)
    axes["battery"].set_title("(d) 储能充放电量")
    axes["battery"].set_ylabel("电量（kWh/10 min）")
    axes["battery"].legend(frameon=False)

    soc_with_initial = np.r_[summary["soc_initial_kwh"], solution.soc]
    axes["soc"].plot(
        interval_edges,
        soc_with_initial,
        color="#0072B2",
        linewidth=2.2,
    )
    axes["soc"].axhline(1200.0, color="#D55E00", linestyle="--", label="下限 1200")
    axes["soc"].axhline(10800.0, color="#009E73", linestyle="--", label="上限 10800")
    axes["soc"].fill_between(
        interval_edges,
        1200.0,
        10800.0,
        color="#0072B2",
        alpha=0.05,
    )
    axes["soc"].set_title("(e) 储能设备电量轨迹")
    axes["soc"].set_ylabel("储电量（kWh）")
    axes["soc"].set_xlabel("时间（h）")
    axes["soc"].legend(frameon=False, ncol=2)

    axes["summary"].axis("off")
    summary_text = (
        "关键结果\n\n"
        f"全天购电量：{summary['total_purchase_q_kwh']:,.2f} kWh\n"
        f"全天购电费：{summary['total_cost_yuan']:,.2f} 元\n"
        f"无储能基线费用：{summary['no_storage_baseline_cost_yuan']:,.2f} 元\n"
        f"相对基线节省：{summary['cost_saving_vs_no_storage_yuan']:,.2f} 元 "
        f"({summary['cost_saving_rate']:.2%})\n"
        f"总充电量 / 放电量：{summary['total_charge_kwh']:,.2f} / "
        f"{summary['total_discharge_kwh']:,.2f} kWh\n"
        f"储电量范围：{summary['soc_min_observed_kwh']:,.2f}–"
        f"{summary['soc_max_observed_kwh']:,.2f} kWh\n"
        f"供电缺口 / 供电富余最大值："
        f"{summary['validation']['max_shortage_kwh']:.2e} / "
        f"{summary['validation']['max_supply_surplus_kwh']:.2e} kWh\n"
        f"状态方程最大残差：{summary['validation']['max_state_residual_kwh']:.2e} kWh"
    )
    axes["summary"].text(
        0.04,
        0.96,
        summary_text,
        va="top",
        fontsize=11,
        linespacing=1.55,
        bbox={"boxstyle": "round,pad=0.7", "facecolor": "#F3F6F8", "edgecolor": "#AAB7C4"},
    )

    for key in ("power", "price", "q", "battery"):
        axes[key].set_xlabel("时间（h）")
    for axis in axes.values():
        if axis.axison:
            axis.set_xlim(0.0, 24.0)
            axis.set_xticks(np.arange(0, 25, 4))
            axis.spines[["top", "right"]].set_visible(False)

    figure.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def plot_key_data(path: Path, summary: dict[str, object]) -> None:
    """生成题目指定时段及四小时充放电汇总表 PNG。"""

    _set_chinese_plot_style()
    figure, axes = plt.subplots(2, 2, figsize=(13, 8), layout="constrained")
    figure.suptitle("问题1关键计算结果", fontsize=18, fontweight="bold")

    axes[0, 0].axis("off")
    axes[0, 0].set_title("全天指标", loc="left", fontweight="bold")
    metrics_text = (
        f"全天购电量：{summary['total_purchase_q_kwh']:,.4f} kWh\n"
        f"全天购电费：{summary['total_cost_yuan']:,.4f} 元\n"
        f"全天负载电量：{summary['total_load_kwh']:,.4f} kWh\n"
        f"光伏可用 / 利用：{summary['total_pv_available_kwh']:,.4f} / "
        f"{summary['total_pv_used_kwh']:,.4f} kWh\n"
        f"初始 / 终止储电量：{summary['soc_initial_kwh']:,.4f} / "
        f"{summary['soc_terminal_kwh']:,.4f} kWh"
    )
    axes[0, 0].text(0.02, 0.88, metrics_text, va="top", fontsize=11.5, linespacing=1.65)

    axes[0, 1].axis("off")
    axes[0, 1].set_title("约束核验", loc="left", fontweight="bold")
    validation = summary["validation"]
    validation_text = (
        f"求解状态：{summary['solver_status']}\n"
        f"最大供电缺口：{validation['max_shortage_kwh']:.3e} kWh\n"
        f"最大供电富余：{validation['max_supply_surplus_kwh']:.3e} kWh\n"
        f"最大光伏越界：{validation['max_pv_excess_kwh']:.3e} kWh\n"
        f"最大充放电越界：{validation['max_throughput_excess_kwh']:.3e} kWh\n"
        f"状态方程最大残差：{validation['max_state_residual_kwh']:.3e} kWh\n"
        f"终止储电量残差：{validation['terminal_soc_residual_kwh']:.3e} kWh"
    )
    axes[0, 1].text(0.02, 0.88, validation_text, va="top", fontsize=11.5, linespacing=1.5)

    axes[1, 0].axis("off")
    axes[1, 0].set_title("题目指定时段的计划购电量", loc="left", fontweight="bold")
    specified_cells = [
        [item["时间段"], f"{item['购电量_kWh']:.4f}"]
        for item in summary["specified_intervals"]
    ]
    specified_table = axes[1, 0].table(
        cellText=specified_cells,
        colLabels=["时间段", "购电量（kWh）"],
        cellLoc="center",
        colLoc="center",
        bbox=[0.02, 0.02, 0.96, 0.88],
    )
    specified_table.auto_set_font_size(False)
    specified_table.set_fontsize(10.5)

    axes[1, 1].axis("off")
    axes[1, 1].set_title("储能设备分时段充放电量", loc="left", fontweight="bold")
    grouped_cells = [
        [item["时间段"], f"{item['充电量_kWh']:.4f}", f"{item['放电量_kWh']:.4f}"]
        for item in summary["four_hour_summary"]
    ]
    grouped_table = axes[1, 1].table(
        cellText=grouped_cells,
        colLabels=["时间段", "充电量（kWh）", "放电量（kWh）"],
        cellLoc="center",
        colLoc="center",
        bbox=[0.0, 0.02, 1.0, 0.88],
    )
    grouped_table.auto_set_font_size(False)
    grouped_table.set_fontsize(10.0)

    for table in (specified_table, grouped_table):
        for (row, _column), cell in table.get_celld().items():
            if row == 0:
                cell.set_facecolor("#1F4E78")
                cell.set_text_props(color="white", weight="bold")
            elif row % 2 == 0:
                cell.set_facecolor("#EEF3F7")

    axes[0, 0].text(
        0.02,
        0.08,
        "说明：数值来自线性规划最优解；充、放电量均为储能接口侧电量。",
        va="bottom",
        fontsize=9,
        color="#555555",
    )
    figure.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(figure)


def run_pipeline(input_path: str | Path, output_dir: str | Path) -> dict[str, object]:
    """读取附件、求解、独立核验并导出全部交付文件。"""

    input_path = Path(input_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    data = load_attachment1(input_path)
    solution = solve_dispatch(data)
    if not solution.success:
        raise RuntimeError(f"线性规划求解失败：{solution.message}")
    validation = validate_solution(data, solution)
    summary = build_summary(data, solution, validation)

    workbook_path = output_dir / "问题1_详细求解结果.xlsx"
    summary_path = output_dir / "问题1_关键指标.json"
    overview_path = output_dir / "问题1_优化结果总览.png"
    key_data_path = output_dir / "问题1_关键数据.png"

    export_workbook(workbook_path, data, solution, summary)
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    plot_overview(overview_path, data, solution, summary)
    plot_key_data(key_data_path, summary)

    return {
        "input": str(input_path),
        "workbook": str(workbook_path),
        "summary_json": str(summary_path),
        "overview_png": str(overview_path),
        "key_data_png": str(key_data_path),
        "summary": summary,
    }


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    local_input = script_dir / "附件1.xlsx"
    default_input = local_input if local_input.exists() else script_dir.parent / "附件1.xlsx"
    parser = argparse.ArgumentParser(description="求解 2026 年国赛 C 题问题 1")
    parser.add_argument(
        "--input",
        type=Path,
        default=default_input,
        help="附件1.xlsx 的路径",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=script_dir,
        help="结果输出文件夹",
    )
    args = parser.parse_args()
    result = run_pipeline(args.input, args.output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
